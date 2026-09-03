from pathlib import Path
from zipfile import ZipFile
import re
from docx import Document
from lxml import etree


DOCX = Path(r"E:\cxy\04-需求管理\PRD-OMS发货单拆单与合单功能-v1.0.docx")
EXPECTED = [
    "1. 需求背景",
    "2. 业务目标",
    "3. 用户痛点",
    "4. 功能范围",
    "5. 优先级说明",
    "6. 原型或流程图",
    "7. 关键规则",
    "8. 异常场景",
    "9. 上线影响",
    "10. 需要协作的依赖项",
]


def main():
    failures = []
    warnings = []
    if not DOCX.exists():
        raise SystemExit(f"FAIL: missing output: {DOCX}")
    if DOCX.stat().st_size < 50_000:
        failures.append(f"file unexpectedly small: {DOCX.stat().st_size} bytes")

    doc = Document(DOCX)
    headings = [p.text.strip() for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for heading in EXPECTED:
        if heading not in headings:
            failures.append(f"missing heading: {heading}")

    text_parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text_parts.append(cell.text)
    full_text = "\n".join(text_parts)

    placeholders = [r"\{\{.+?\}\}", r"\[\[.+?\]\]", r"\bTODO\b", r"\bTBD\b"]
    for pattern in placeholders:
        matches = re.findall(pattern, full_text, flags=re.I)
        if matches:
            failures.append(f"placeholder tokens found for {pattern}: {matches[:3]}")

    required_terms = [
        "旺店通驳回订单",
        "发货管理",
        "数量守恒",
        "已处理",
        "乐观校验",
        "拆/合单管理",
        "拆单记录",
        "合单记录",
        "处理后订单",
        "处理前订单",
        "二次驳回",
    ]
    for term in required_terms:
        if term not in full_text:
            failures.append(f"missing required term: {term}")

    if "4.3 本期不包含" in full_text:
        failures.append("removed section still exists: 4.3 本期不包含")

    fake_bullets = [p.text for p in doc.paragraphs if p.text.strip().startswith(("•", "- "))]
    if fake_bullets:
        warnings.append(f"possible fake bullets: {fake_bullets[:3]}")

    with ZipFile(DOCX) as zf:
        names = set(zf.namelist())
        media = [name for name in names if name.startswith("word/media/")]
        if len(media) < 2:
            failures.append(f"expected at least 2 diagram images, found {len(media)}")
        document_xml = etree.fromstring(zf.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        tables = document_xml.xpath("//w:tbl", namespaces=ns)
        for index, table in enumerate(tables, 1):
            grid_widths = [int(x.get("{%(w)s}w" % ns)) for x in table.xpath("./w:tblGrid/w:gridCol", namespaces=ns)]
            if grid_widths and sum(grid_widths) != 9360:
                warnings.append(f"table {index} grid total {sum(grid_widths)} DXA (expected 9360)")
            tbl_w = table.xpath("./w:tblPr/w:tblW", namespaces=ns)
            if tbl_w:
                width = tbl_w[0].get("{%(w)s}w" % ns)
                if width not in ("9360", "0"):
                    warnings.append(f"table {index} tblW={width}")

    print(f"DOCX={DOCX}")
    print(f"SIZE={DOCX.stat().st_size} bytes")
    print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} HEADINGS={len(headings)} IMAGES={len(media)}")
    print(f"TEXT_LENGTH={len(full_text)}")
    print("HEADINGS:")
    for heading in headings:
        print(f"  - {heading}")
    for warning in warnings:
        print(f"WARN: {warning}")
    if failures:
        for failure in failures:
            print(f"FAIL: {failure}")
        raise SystemExit(1)
    print("PASS: structural PRD audit completed")


if __name__ == "__main__":
    main()
