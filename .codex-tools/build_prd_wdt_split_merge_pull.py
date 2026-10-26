from pathlib import Path
import importlib.util

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"E:\cxy")
OUT_FILE = ROOT / "04-需求管理" / "PRD-OMS拉取旺店通拆合单数据-v1.0.docx"
ASSET_DIR = ROOT / ".codex-tools" / "prd_wdt_split_merge_pull_assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

spec = importlib.util.spec_from_file_location("prd_base", ROOT / ".codex-tools" / "build_prd_split_merge.py")
base = importlib.util.module_from_spec(spec)
spec.loader.exec_module(base)


def font(bold=False, size=24):
    return ImageFont.truetype(base.font_path(bold), size)


def make_business_flow(path: Path):
    img = Image.new("RGB", (1800, 760), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 38), "发货单同步与旺店通拆合记录回传流程", fill="#173A5E", font=font(True, 38))
    d.text((60, 92), "OMS主发货单始终保持一条，旺店通拆合结果只作为关联历史。", fill="#5D6B78", font=font(False, 22))
    boxes = [
        ((45, 205, 305, 350), "OMS原始发货单", "列表保留1条主记录"),
        ((370, 205, 630, 350), "同步至旺店通", "发送1条发货单数据"),
        ((695, 205, 955, 350), "旺店通拆单或合单", "生成来源单与结果单关系"),
        ((1020, 205, 1280, 350), "OMS增量拉取", "幂等写入拆合历史"),
        ((1345, 205, 1755, 350), "发货管理展示", "仍为1条发货单 + 查看拆合记录"),
    ]
    colors = [
        ("#EAF4FD", "#77B5E8"), ("#EAF4FD", "#77B5E8"),
        ("#FFF3E8", "#E8A55A"), ("#EAF6EC", "#77B980"),
        ("#EAF4FD", "#77B5E8"),
    ]
    for (box, title, subtitle), (fill, outline) in zip(boxes, colors):
        base.draw_rounded_box(d, box, fill, outline, title, subtitle)
    for left, right in zip(boxes, boxes[1:]):
        base.arrow(d, (left[0][2] + 8, 278), (right[0][0] - 8, 278))

    d.rounded_rectangle((250, 500, 760, 640), radius=16, fill="#EAF6EC", outline="#77B980", width=2)
    d.text((505, 538), "匹配成功", fill="#18324B", font=font(True, 29), anchor="mm")
    d.text((505, 590), "挂载到原发货单的拆合记录", fill="#506070", font=font(False, 22), anchor="mm")
    d.rounded_rectangle((1040, 500, 1550, 640), radius=16, fill="#FFF3E8", outline="#E8A55A", width=2)
    d.text((1295, 538), "匹配失败或数量异常", fill="#18324B", font=font(True, 29), anchor="mm")
    d.text((1295, 590), "后台保留异常，不新增发货单", fill="#506070", font=font(False, 22), anchor="mm")
    base.arrow(d, (1150, 360), (760, 500))
    base.arrow(d, (1150, 360), (1040, 500))
    img.save(path)


def make_display_model(path: Path):
    img = Image.new("RGB", (1800, 760), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 38), "发货管理与拆合单管理双入口模型", fill="#173A5E", font=font(True, 38))
    d.text((60, 92), "同一套旺店通关系数据，在两个页面按不同业务目的展示。", fill="#5D6B78", font=font(False, 22))

    base.draw_rounded_box(d, (65, 180, 525, 330), "#EAF4FD", "#77B5E8", "发货管理", "OMS发货单F001：始终1条")
    base.draw_rounded_box(d, (65, 470, 525, 620), "#FFF3E8", "#E8A55A", "拆/合单管理", "按拆合记录展示被处理订单")
    d.rounded_rectangle((720, 130, 1725, 665), radius=18, fill="#FAFBFC", outline="#B5C0CC", width=2)
    d.text((765, 175), "旺店通拆合关系数据", fill="#18324B", font=font(True, 29))
    base.draw_rounded_box(d, (785, 245, 1110, 390), "#FFF3E8", "#E8A55A", "处理前订单", "来源单A：商品3盒")
    base.draw_rounded_box(d, (1335, 245, 1660, 390), "#EAF6EC", "#77B980", "处理后订单", "结果A1：1盒 / A2：2盒")
    base.arrow(d, (1120, 318), (1325, 318))
    base.draw_rounded_box(d, (785, 485, 1110, 610), "#F4F6F9", "#B5C0CC", "关联OMS主记录", "F001 / 原订单 / 原商品")
    base.draw_rounded_box(d, (1335, 485, 1660, 610), "#EAF4FD", "#77B5E8", "查看方式", "主单详情 + 集中管理")
    base.arrow(d, (1120, 547), (1325, 547))
    base.arrow(d, (535, 255), (710, 255))
    base.arrow(d, (535, 545), (710, 545))
    img.save(path)


def set_document_styles(doc: Document):
    base.setup_styles(doc)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = base.rgb("000000")
    doc.styles["Title"].font.size = Pt(25)
    doc.styles["Title"].font.bold = True


def add_metadata(doc, rows):
    table = base.add_table(doc, ["项目属性", "内容"], rows, [1900, 7460], header_fill=base.LIGHT_BLUE, font_size=9.5)
    table.alignment = base.WD_TABLE_ALIGNMENT.LEFT


def build_doc():
    business_flow = ASSET_DIR / "wdt-pull-flow.png"
    display_model = ASSET_DIR / "wdt-relation-flow.png"
    make_business_flow(business_flow)
    make_display_model(display_model)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.footer_distance = Inches(0.35)
    set_document_styles(doc)
    base.add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("OMS发货管理与旺店通拆合单管理")
    base.set_run_font(run, size=25, color="000000", bold=True)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    run = sub.add_run("产品需求文档")
    base.set_run_font(run, size=13, color=base.MUTED)

    add_metadata(doc, [
        ["文档版本", "v1.0"], ["文档状态", "评审稿"],
        ["所属系统", "OMS订单系统 / 发货管理"], ["关联平台", "旺店通"],
        ["需求负责人", "陈薪宇"], ["编制日期", "2026-10-26"],
        ["原型文件", "旺店通拆合单数据拉取交互演示.html"],
    ])
    base.add_para(doc, "产品定义：OMS发货管理继续以原始发货单为唯一主记录；拆/合单管理独立集中展示旺店通返回的被拆合订单数据。两个页面读取同一套拆合关系：前者保证发货主单稳定，后者提供处理前订单、处理后订单和商品数量的集中查询。", bold_lead="产品定义：", size=10.8, after=12)

    base.add_heading(doc, "1 需求背景", 1)
    base.add_para(doc, "OMS现有发货流程以发货单为主记录，一张发货单同步至旺店通后，旺店通可能根据履约需要执行拆单或合单。旺店通拆合后会产生新的来源单与结果单关系，但这些单号属于旺店通执行层，不应反向改变OMS的发货单主数据。")
    base.add_para(doc, "发货管理与拆合业务查询存在不同目标：发货管理需要维持OMS原始发货单口径，拆合业务人员则需要集中查看所有被旺店通拆单或合单的订单。因此系统保留两个入口：发货管理中从原发货单查看关联记录，拆/合单管理中按拆合事件查看处理前后订单数据。")
    base.add_para(doc, "旺店通是拆合结果的权威数据源。OMS不执行拆合、不修改旺店通结果，也不根据旺店通结果生成新的OMS发货单。无法匹配的返回数据进入后台异常处理，不能污染发货管理主列表。")

    base.add_heading(doc, "2 业务目标", 1)
    base.add_bullets(doc, [
        ("主记录稳定：", "无论旺店通拆分为几张单或合并多少来源单，OMS原始发货单在列表中始终只保留一条。"),
        ("处理可追溯：", "从原发货单直接查看旺店通拆单、合单的来源单、结果单、商品及数量。"),
        ("集中可查询：", "通过拆/合单管理统一查看全部被旺店通拆合的订单数据。"),
        ("同步边界清晰：", "OMS只负责首次同步发货单和后续拉取拆合结果，不在OMS内执行拆合。"),
        ("异常不污染主数据：", "未匹配、数量异常或接口失败时保留原始返回并告警，不新增孤立发货单。"),
        ("下游查询一致：", "物流、售后和退货从原发货单沿拆合记录定位旺店通实际执行单。"),
    ])
    base.add_para(doc, "成功标准：同一OMS发货单在发货管理列表中始终只有一条；拆合数据拉取成功后，可从原发货单查看处理历史，也可在拆/合单管理中查询全部被拆合订单，并能从结果单回溯原始发货单。", bold_lead="成功标准：")

    base.add_heading(doc, "3 用户痛点", 1)
    base.add_table(doc, ["用户角色", "当前问题", "目标体验"], [
        ["发货运营人员", "旺店通拆合后出现多个执行单号，容易误认为OMS也应新增多条发货单。", "发货管理保持原始一条记录，按需查看拆合历史。"],
        ["客服与售后人员", "仅知道旺店通结果单时，难以定位最初OMS发货单。", "从任一结果单追溯原发货单、原商品和数量。"],
        ["技术支持人员", "重复拉取或字段缺失可能产生重复主记录或错误关联。", "幂等挂载关系；异常数据与主列表隔离。"],
        ["业务负责人", "仅在单张发货单内查看记录，不便集中掌握所有被拆合订单。", "在拆/合单管理集中查看，同时保持发货主单口径不变。"],
    ], [1550, 3900, 3910], font_size=8.9)

    base.add_heading(doc, "4 功能范围", 1, page_break=True)
    base.add_heading(doc, "4.1 本期功能", 2)
    base.add_table(doc, ["模块", "功能", "页面或任务"], [
        ["发货单首次同步", "继续按现有逻辑将一条OMS发货单同步为一条旺店通原始单，并保存单号映射。", "发货管理 / 同步接口"],
        ["拆合数据拉取", "定时增量拉取旺店通已完成的拆单、合单记录；支持人工更新旺店通状态。", "后台任务 / 发货管理按钮"],
        ["原发货单挂载", "将匹配成功的拆合记录挂载到原始OMS发货单，不新增OMS发货单。", "关系服务"],
        ["列表摘要", "原发货单行展示是否存在拆合记录、记录数、当前旺店通结果单和最近更新时间。", "发货管理列表"],
        ["拆合记录详情", "展示拆合时间线、来源单、结果单、处理前后商品数量和合单关联的其他来源单。", "详情弹窗"],
        ["拆/合单管理", "按拆合事件集中展示被处理的订单；支持全部、拆单、合单Tab及条件查询。", "独立菜单页面"],
        ["拆合订单卡片", "统一按处理前订单在左、处理后订单在右展示旺店通单号、关联OMS单据及商品数量。", "拆/合单管理列表"],
        ["查询能力", "发货管理支持按原始或结果旺店通单号查询并定位原发货单。", "筛选区"],
        ["后台异常处理", "未匹配、映射冲突和数量异常数据保留原始响应并支持技术重试。", "后台任务与监控"],
    ], [1750, 5000, 2610], font_size=8.7)

    base.add_heading(doc, "4.2 本期不包含", 2)
    base.add_bullets(doc, [
        "拆/合单管理不承载未匹配原始数据或技术异常处理，异常数据继续在后台任务与监控中处理。",
        "不把旺店通拆分结果单或合并结果单转换为新的OMS发货单。",
        "不在OMS内执行拆单、合单、商品数量分配或合单候选选择。",
        "不允许业务用户修改旺店通来源单、结果单或商品数量。",
        "不增加仓库、快递、包裹、运费等拆合附加信息。",
        "不根据客户、地址或商品名称进行高风险模糊自动匹配。",
    ])

    base.add_heading(doc, "5 优先级说明", 1)
    base.add_table(doc, ["优先级", "范围", "验收重点"], [
        ["P0", "OMS发货单唯一主记录", "旺店通拆合前后，发货管理记录数和OMS发货单号不变。"],
        ["P0", "首次同步与单号映射", "一条OMS发货单只同步一次并保存旺店通原单映射。"],
        ["P0", "拆合记录拉取与幂等", "重复拉取只更新同一历史记录，不生成重复关系或主记录。"],
        ["P0", "列表摘要与详情入口", "有拆合记录时可见记录数，并能打开完整详情。"],
        ["P0", "拆/合单管理", "可集中筛选拆单与合单记录，并展示处理前订单和处理后订单。"],
        ["P0", "来源结果与数量展示", "拆单展示处理前和各结果数量；合单展示各来源及合并结果数量。"],
        ["P0", "异常隔离", "匹配失败或数量异常时不在发货管理新增记录。"],
        ["P1", "跨单号查询", "输入旺店通原单或结果单均可定位原OMS发货单。"],
        ["P2", "多级关系链", "展示结果单再次拆合后的完整处理时间线。"],
    ], [900, 2850, 5610], font_size=8.7)

    base.add_heading(doc, "6 原型或流程图", 1, page_break=True)
    base.add_heading(doc, "6.1 业务流程", 2)
    doc.add_picture(str(business_flow), width=Inches(6.65))
    doc.inline_shapes[-1]._inline.docPr.set("title", "发货单同步与旺店通拆合记录回传流程")
    doc.inline_shapes[-1]._inline.docPr.set("descr", "OMS保留一条发货单并同步旺店通，旺店通拆合后由OMS拉取历史并挂载到原发货单。")
    base.add_caption(doc, "图1 发货单同步与旺店通拆合记录回传流程")
    base.add_heading(doc, "6.2 页面展示关系", 2)
    doc.add_picture(str(display_model), width=Inches(6.65))
    doc.inline_shapes[-1]._inline.docPr.set("title", "发货管理与拆合单管理双入口模型")
    doc.inline_shapes[-1]._inline.docPr.set("descr", "发货管理保持一条OMS主发货单，拆合单管理集中展示旺店通处理前后订单，两者共享关系数据。")
    base.add_caption(doc, "图2 发货管理与拆合单管理双入口模型")
    base.add_heading(doc, "6.3 页面原型", 2)
    base.add_table(doc, ["页面区域", "核心内容"], [
        ["发货管理列表", "沿用现有发货单卡片结构，不增加拆合结果行；一张原发货单只展示一次。"],
        ["筛选区", "保留发货单、订单、客户、同步状态等条件，并支持按旺店通原单或结果单号定位。"],
        ["拆合摘要", "在原发货单行展示拆合类型、记录数、当前结果单数量和最近更新时间。"],
        ["查看拆合单记录", "弹窗顶部先展示原始OMS发货单，再按时间顺序展示旺店通拆单与合单历史。"],
        ["拆/合单管理", "独立菜单集中展示全部被拆合订单，提供全部记录、拆单记录、合单记录Tab。"],
        ["管理列表卡片", "每条记录先展示处理前订单，再展示处理后订单，并标注关联OMS发货单和原订单。"],
        ["拆单记录", "先展示处理前旺店通单和原商品数量，再展示各结果单及分配数量。"],
        ["合单记录", "先展示各来源旺店通单和来源商品数量，再展示合并结果单；说明相关OMS发货单仍各自保留。"],
    ], [2200, 7160], font_size=9.0)
    base.add_para(doc, r"交互原型路径：E:\cxy\09-设计与原型\旺店通拆合单数据拉取交互演示.html", color=base.MUTED, size=9.2)

    base.add_heading(doc, "7 关键规则", 1, page_break=True)
    base.add_heading(doc, "7.1 OMS发货单主记录规则", 2)
    base.add_table(doc, ["规则项", "要求"], [
        ["唯一主记录", "OMS发货单是发货管理唯一主记录。旺店通拆单、合单及多级拆合均不改变其编号和列表记录数。"],
        ["原始业务数据", "OMS保留原订单、原商品行和原发货数量，不使用旺店通结果覆盖主记录。"],
        ["拆合从属数据", "旺店通拆合记录以子表或关系表挂载到原发货单，通过详情入口读取。"],
        ["合单多主关联", "一条旺店通合单记录可关联多张OMS原发货单；各原发货单仍各自展示，并可查看同一合单历史。"],
        ["列表展示", "默认不展开拆合明细，只展示有无记录、记录数、当前结果单和查看入口。"],
        ["删除约束", "原发货单被业务引用后不得因旺店通关系更新而删除；拆合历史保留版本。"],
    ], [2200, 7160], font_size=8.8)

    base.add_heading(doc, "7.2 双入口展示规则", 2)
    base.add_table(doc, ["页面", "展示口径"], [
        ["发货管理", "一张OMS原发货单只展示一条；拆合信息作为摘要和详情挂载，不增加结果行。"],
        ["拆/合单管理", "一条旺店通拆合记录展示一条；集中展示被处理订单，不参与OMS发货单数量统计。"],
        ["拆单记录", "先展示一张处理前来源单及原商品数量，再展示多张处理后结果单及分配数量。"],
        ["合单记录", "先展示多张处理前来源单及各自商品数量，再展示一张处理后合并结果单。"],
        ["跨页追溯", "从发货管理可进入该发货单记录；从拆/合单管理可反查所有关联OMS发货单。"],
        ["数据一致性", "两个页面共享同一关系ID和更新时间，不复制生成两套业务记录。"],
    ], [2200, 7160], font_size=8.8)

    base.add_heading(doc, "7.3 数据拉取与匹配规则", 2)
    base.add_table(doc, ["规则项", "要求"], [
        ["权威数据源", "旺店通拆合记录是执行结果的唯一权威来源，OMS不得自行推导或改写。"],
        ["触发方式", "后台定时增量拉取；发货管理可提供人工更新旺店通状态。两者共用游标和幂等规则。"],
        ["增量游标", "使用旺店通更新时间与记录唯一ID组合推进水位，避免同一时间多条记录遗漏。"],
        ["主记录匹配", "优先使用OMS发货单唯一ID、原同步请求号或既有旺店通原单映射定位原发货单。"],
        ["幂等写入", "同一旺店通拆合记录重复返回时更新原历史，不重复挂载。"],
        ["匹配失败", "保存原始数据并进入后台异常处理；不在发货管理创建临时或孤立发货单。"],
        ["只读边界", "OMS拉取任务不向旺店通写入拆合数据，也不触发重新拆合。"],
    ], [2200, 7160], font_size=8.8)

    base.add_heading(doc, "7.4 关系与数量规则", 2)
    base.add_numbered(doc, [
        "拆单记录保存一张来源单、拆单记录ID、多张结果单，以及每个商品行的处理前数量和各结果数量。",
        "合单记录保存多张来源单、合单记录ID、一张结果单，并按OMS原发货单商品行保留来源数量。",
        "拆单结果数量合计必须等于来源数量；合单结果数量必须等于全部来源数量合计。",
        "结果单再次参与拆单或合单时新增下一条历史，按旺店通操作时间排序，不覆盖上一条记录。",
        "当前结果单是关系链中没有后续拆合的旺店通结果单；它只用于执行追溯，不替代OMS发货单号。",
    ])

    base.add_heading(doc, "7.5 状态与权限", 2)
    base.add_table(doc, ["对象", "状态或权限"], [
        ["OMS发货单", "沿用现有发货、同步、接单等状态，不新增拆单后发货单状态。"],
        ["拆合摘要", "无拆合记录、有拆合记录、数据异常；状态同时使用文字和颜色。"],
        ["人工更新", "需要旺店通状态更新权限；执行中禁用按钮，成功或失败给出明确反馈。"],
        ["拆合详情", "具备发货管理查询权限的人员可查看，但不能编辑旺店通记录和数量。"],
        ["拆/合单管理", "具备拆合记录查询权限的人员可查看和导出，不提供编辑、删除或重新拆合。"],
    ], [2200, 7160], font_size=8.8)

    base.add_heading(doc, "8 异常场景", 1, page_break=True)
    base.add_table(doc, ["异常场景", "系统处理", "页面反馈"], [
        ["旺店通接口超时或不可用", "本次更新失败或进入重试，不推进成功水位。", "保留原发货单及上次成功记录，提示失败原因。"],
        ["重复返回同一拆合记录", "按旺店通记录ID幂等更新。", "拆合记录数不重复增加。"],
        ["旺店通记录无法匹配原发货单", "保存原始返回并进入后台异常处理。", "发货管理不新增孤立记录；必要时显示数据待更新。"],
        ["多个OMS发货单同时命中", "停止自动绑定并记录候选冲突。", "不覆盖现有关系，由技术核对稳定标识。"],
        ["拆单数量不守恒", "记录标记数据异常，不作为有效执行关系供下游使用。", "详情显示处理前数量、结果合计和差额。"],
        ["合单缺少部分来源单", "保留已返回信息但不标记关系完整。", "详情提示来源数据不完整。"],
        ["返回顺序错乱", "先保存原始记录，待原单映射到达后自动重试匹配。", "不生成新的发货单，最近更新时间保持可见。"],
        ["拆合关系被旺店通更新", "按更新时间更新当前版本并保留历史快照。", "详情展示最新结果和更新时间。"],
        ["人工更新与定时任务并发", "复用进行中任务或通过互斥锁避免重复推进游标。", "按钮显示更新中或提示稍后重试。"],
        ["合单关联其他OMS发货单", "同一合单记录分别挂载到所有已匹配原发货单。", "当前详情列出其他来源单，但列表仍各自保留原记录。"],
    ], [2500, 4450, 2410], font_size=8.2)

    base.add_heading(doc, "9 上线影响", 1, page_break=True)
    base.add_heading(doc, "9.1 业务影响", 2)
    base.add_bullets(doc, [
        "发货人员继续按现有方式在OMS生成一张发货单并同步旺店通。",
        "拆单和合单继续在旺店通执行，OMS发货管理不新增拆合操作。",
        "旺店通处理完成后，业务人员回到原发货单点击“查看拆合单记录”追溯执行结果。",
        "需要集中核对时，业务人员进入拆/合单管理查询所有被拆合订单和处理结果。",
        "发货管理列表、统计和导出口径继续按OMS原发货单计算，不能按旺店通结果单增加记录数。",
        "客服、售后和退货可从原发货单沿拆合记录定位实际旺店通单号和商品数量。",
    ])
    base.add_heading(doc, "9.2 系统影响", 2)
    base.add_table(doc, ["影响面", "改造内容"], [
        ["数据模型", "在现有发货单外新增拆合记录、来源单、结果单、商品行数量、原始响应和匹配结果关系表。"],
        ["接口任务", "新增定时增量拉取、人工更新、分页重试、水位管理和限流处理。"],
        ["发货管理列表", "不新增拆合结果行；增加拆合摘要、旺店通单号查询和查看记录入口。"],
        ["发货单详情", "增加只读拆合时间线，先展示原记录，再展示旺店通处理结果。"],
        ["拆/合单管理", "新增独立菜单、筛选、拆单与合单Tab、处理前后订单卡片和详情入口。"],
        ["统计与导出", "仍按OMS发货单主键去重统计；拆合明细仅在需要时单独导出。"],
        ["物流与售后", "从原发货单关系链读取当前旺店通结果单；异常关系不得进入自动处理。"],
        ["监控", "增加拉取失败、未匹配、数量异常、映射冲突和关系积压告警。"],
    ], [2300, 7060], font_size=8.9)
    base.add_heading(doc, "9.3 上线建议", 2)
    base.add_numbered(doc, [
        "确认旺店通拆合接口字段，并准备拆单、合单、先拆后合和多发货单合并样例。",
        "先上线关系表和拉取任务，验证不会改变现有发货单记录数、编号和状态。",
        "灰度开放发货管理的拆合摘要与详情入口，抽查旺店通数量和单号关系。",
        "确认统计、导出、物流、售后和退货均以OMS原发货单为主口径后再全量发布。",
        "回滚时关闭新的关系展示和拉取任务，保留已拉取原始数据，不删除原发货单。",
    ])

    base.add_heading(doc, "10 需要协作的依赖项", 1, page_break=True)
    base.add_table(doc, ["协作方", "依赖事项", "需要确认或交付", "优先级"], [
        ["旺店通接口负责人", "拆合记录查询接口", "提供记录ID、更新时间、来源单、结果单、商品行、数量及可反查OMS原发货单的稳定标识。", "P0"],
        ["OMS后端", "关系模型与拉取任务", "实现游标、幂等、分页重试、主记录匹配、历史挂载和异常隔离。", "P0"],
        ["OMS前端", "双入口页面改造", "发货管理保持一条主记录并增加详情；拆/合单管理实现筛选、Tab、处理前后订单卡片。", "P0"],
        ["订单与发货服务", "稳定主键与统计口径", "保证拆合结果不生成新发货单，并提供原同步请求号和旺店通原单映射。", "P0"],
        ["物流与售后系统", "结果单读取规则", "从OMS原发货单沿关系链读取当前旺店通执行单，并定义异常阻断规则。", "P0"],
        ["业务运营", "页面验收样例", "确认一张发货单拆分、跨单合并和先拆后合的展示方式。", "P0"],
        ["测试团队", "端到端测试", "覆盖记录数不变、重复拉取、跨单查询、乱序、未匹配、数量异常和并发任务。", "P0"],
        ["运维与监控", "调度与告警", "配置任务执行、接口失败、未匹配积压和数据异常告警。", "P1"],
    ], [1550, 2050, 4860, 900], font_size=8.1)

    base.add_heading(doc, "10.1 联调前确认清单", 2)
    base.add_bullets(doc, [
        "旺店通拆合记录是否能返回OMS发货单唯一ID、原同步请求号或可稳定反查的原单号。",
        "拆单和合单记录的唯一ID、更新时间精度、分页方式和同一时间多条数据的排序字段。",
        "合单涉及多张OMS发货单时，接口是否返回全部来源单和各来源商品数量。",
        "结果单再次拆合时，接口如何表达上游记录和当前最终结果单。",
        "原发货单统计、导出、物流、售后和退货是否统一以OMS发货单主键去重。",
        "历史数据初始化范围，以及上线前后发货单记录数不变的核对方式。",
    ])

    props = doc.core_properties
    props.title = "OMS发货管理与旺店通拆合单管理 PRD"
    props.subject = "原始发货单唯一主记录与被拆合订单集中查询"
    props.author = "卓希产品团队"
    props.keywords = "OMS, 旺店通, 发货管理, 拆单记录, 合单记录, 数据拉取, PRD"
    doc.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    build_doc()
