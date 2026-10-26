from pathlib import Path
import zipfile

from docx import Document


path = Path(r"E:\cxy\04-需求管理\PRD-OMS拉取旺店通拆合单数据-v1.0.docx")
doc = Document(path)
text = "\n".join(
    [p.text for p in doc.paragraphs]
    + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
)
required = [
    "1 需求背景",
    "2 业务目标",
    "3 用户痛点",
    "4 功能范围",
    "5 优先级说明",
    "6 原型或流程图",
    "7 关键规则",
    "8 异常场景",
    "9 上线影响",
    "10 需要协作的依赖项",
    "不在OMS内执行拆单、合单",
    "旺店通拆合记录是拆合结果的唯一权威来源",
    "待匹配池",
]
missing = [item for item in required if item not in text]
assert not missing, f"missing: {missing}"
forbidden = [
    "确认拆单并生成子单",
    "本期按P0范围建设“OMS当前有效发货单处理”能力",
]
found = [item for item in forbidden if item in text]
assert not found, f"forbidden: {found}"
with zipfile.ZipFile(path) as archive:
    media = [name for name in archive.namelist() if name.startswith("word/media/")]
assert len(media) == 2, media
headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
assert headings[:10] == [
    "1 需求背景",
    "2 业务目标",
    "3 用户痛点",
    "4 功能范围",
    "4.1 本期功能",
    "4.2 本期不包含",
    "5 优先级说明",
    "6 原型或流程图",
    "6.1 数据拉取流程",
    "6.2 拆单与合单关系",
], headings[:10]
print("PRD_WDT_PULL_VALIDATION_OK")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} media={len(media)} bytes={path.stat().st_size}")
