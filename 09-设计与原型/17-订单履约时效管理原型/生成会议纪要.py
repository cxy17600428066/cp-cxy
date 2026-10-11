from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).with_name('订单履约时效调整会议纪要.docx')
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)

def set_font(run, name='微软雅黑', size=None, bold=None, color='000000'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

styles = doc.styles
normal = styles['Normal']
normal.font.name = '微软雅黑'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.28

for sty_name, size, before, after in [('Title', 22, 0, 18), ('Heading 1', 15, 16, 8), ('Heading 2', 12, 10, 5)]:
    st = styles[sty_name]
    st.font.name = '微软雅黑'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    pPr = st._element.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is not None: pPr.remove(pBdr)

if '正文要点' not in styles:
    s = styles.add_style('正文要点', WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = styles['Normal']
    s.paragraph_format.left_indent = Cm(0.65)
    s.paragraph_format.first_line_indent = Cm(-0.35)
    s.paragraph_format.space_after = Pt(2.5)

def add_title(text):
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is not None: pPr.remove(pBdr)
    set_font(p.add_run(text), size=22, bold=True)

def add_intro(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    set_font(p.add_run(text), size=11)

def h1(text): doc.add_paragraph(text, style='Heading 1')
def h2(text): doc.add_paragraph(text, style='Heading 2')

def body(text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True)
        set_font(p.add_run(text[len(bold_lead):]))
    else:
        set_font(p.add_run(text))
    return p

def bullet(text, level=0):
    p = doc.add_paragraph(style='正文要点')
    p.paragraph_format.left_indent = Cm(0.7 + level * 0.55)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    set_font(p.add_run('• ' + text))
    return p

def number_item(number, text):
    p = doc.add_paragraph(style='正文要点')
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    set_font(p.add_run(f'{number}. '), bold=True)
    set_font(p.add_run(text))
    return p

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc.get_or_add_tcPr()
    tcMar = tc.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tc.append(tcMar)
    for tag, val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+tag))
        if node is None: node=OxmlElement('w:'+tag); tcMar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    t.autofit = False
    for i, head in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,'1F4E78'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(head), size=9.5, bold=True, color='FFFFFF')
        if widths: c.width=Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells=t.add_row().cells
        for i, val in enumerate(row):
            c=cells[i]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
            if ridx%2: shade(c,'F3F7FA')
            p=c.paragraphs[0]
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(str(val)), size=9.5)
            if widths: c.width=Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

add_title('订单履约时效调整会议纪要')
add_intro('本次会议围绕客户分级、下单时效选择、工厂接单、库存校验、重点客户关注及物流资料采集进行梳理。后续系统应以真实业务事件为依据，明确客户要求时间、预计时间和实际完成时间，并通过 OMS、OA、企业微信、库存及物流系统完成协同。')

h1('一 客户分级 签约主体与到货时效配置')
body('客户按照组织层级分为三级，并按各级收货信息分别维护到货时效。不同公司、仓库和门店可以独立配置，不统一沿用集团层级的时效。')
table(['客户层级','对象','配置要求'],[
    ['一级','集团公司','维护集团级客户信息、签约关系及通用时效规则'],
    ['二级','公司、仓','按公司或仓库维护收货信息和到货时效'],
    ['三级','门店','按具体门店维护收货地址、联系人及到货时效'],
],[2.2,3.2,10.0])
body('同时需要确认客户对应的签约主体，并根据签约主体及约定的收费标准收取相应费用。系统应记录以下信息：')
for x in ['客户组织层级','签约主体','费用项目和收费标准','费用承担方','规则生效及失效时间','合同或协议附件']:
    bullet(x)
body('下单时，系统根据客户、收货主体和签约主体自动匹配到货时效及收费规则；无法匹配或存在多条有效规则时，提示业务人员确认后再提交订单。')

h1('二 下单时的时效选择')
h2('经销客户')
bullet('默认选择“工厂默认时间”。')
bullet('允许改选“指定时间”。选择指定时间时，必须上传聊天记录或其他证明附件，作为客户时间要求的依据。')
h2('大客户')
bullet('默认选择“指定时间”。')
bullet('指定时间来源于客户要求及对应的客户时效配置。')

h1('三 大客户接单流程')
body('大客户接单分为“标准接单”和“快速接单”。')
h2('标准接单')
body('按期发货：', '按期发货：')
bullet('工厂确认可以按期履约。')
bullet('根据商品库存情况，分别执行已配置的“有货发货时效”或“无货发货时效”。')
bullet('系统按照对应节点持续监控履约进度。')
body('申请延期：', '申请延期：')
bullet('工厂填写延期时长、延期后的计划时间及延期原因。')
bullet('延期申请对接 OA，由张勇总审核，并通过企业微信发送提醒。')
bullet('审批通过后，系统调整发货时间并按新时间监控。')
bullet('审批未通过时，退回工厂重新选择“按期发货”或再次提交延期方案。')
h2('快速接单')
bullet('支持自定义配置监控节点。')
bullet('系统按照所选节点及对应时效进行履约监控。')
bullet('具体可配置节点和适用条件需要进一步确认。')

h1('四 经销商接单流程')
body('经销商订单根据下单时选择的时效类型，分为“工厂默认时间发货”和“指定时间发货”。')
h2('工厂默认时间发货')
body('有货发货：', '有货发货：')
bullet('按照已配置的有货发货时效执行，并持续监控对应履约节点。')
body('无货发货：', '无货发货：')
steps=[
 '工厂接单时选择最晚发货时间。系统可根据配置提供默认值，工厂可以修改。',
 '修改后，将发货时间同步给下单员确认，并通过企业微信通知下单员进入 OMS 处理。',
 '双方确认时间后，进入后续履约流程。',
 '下单员认为时间无法接受时，可以拒绝并提出新的发货时间，由工厂再次确认。',
 '工厂接受下单员调整后的时间后，进入后续履约流程。',
 '工厂无法接受下单员调整时间时，可以发起 OA 协调申请，由佳佳协调并上传双方确认后的时效。',
 '佳佳提交协调结果后，如果工厂仍无法满足要求，继续发起 OA 审批，由严总、张勇总审批。',
 '审批完成后，由佳佳回传最终时间及处理结果，系统按照最终结果更新并监控履约计划。'
]
for i,x in enumerate(steps,1): number_item(i,x)
h2('指定时间发货')
body('按期发货：', '按期发货：')
bullet('工厂确认可以满足客户指定时间。')
bullet('根据商品库存情况，执行已配置的有货或无货发货时效。')
bullet('系统按照指定到货时间及各履约节点进行监控。')
body('申请延期：', '申请延期：')
bullet('工厂填写延期时长、延期后的计划时间及延期原因。')
bullet('延期申请对接 OA，由张勇总审批，并通过企业微信发送提醒。')
bullet('审批通过后更新发货时间；审批未通过时退回工厂重新选择处理方式。')

h1('五 接单时的商品库存校验')
body('接单页面需要展示商品实时库存信息：')
for x in ['库存数量','已占用数量','可使用数量']:
    bullet(x)
bullet('可使用库存足以满足本次订单时，不允许选择“无货发货”。')
bullet('库存不足时，可以选择“无货发货”，并进入相应的时间确认或协调流程。')
bullet('多商品订单逐项展示库存情况，并明确提示影响接单选择的商品。')

h1('六 重点客户订单关注')
body('针对盒马等重点大客户，增加“重点客户订单关注”功能，采用总分结构展示。')
h2('汇总层')
for x in ['本日下单数量','订单金额','待接单数量','待发货数量','异常或即将超时订单数量']:
    bullet(x)
h2('明细层')
for x in ['订单编号和下单时间','客户及收货主体','商品明细','当前履约节点和订单状态','预计时间与实际完成时间','异常或超时情况']:
    bullet(x)
body('统计数据必须来源于真实订单和履约记录，缺少数据时显示“待同步”或“暂无数据”。')

h1('七 专车物流信息采集')
body('对于未对接物流平台、由业务人员自行安排专车配送的订单，需要采集以下资料，并与对应订单、发货单及运输任务关联：')
for x in ['司机姓名','手机号','车牌号','驾驶证','行驶证']:
    bullet(x)
body('驾驶证和行驶证支持上传图片或文件。')

h1('八 自提订单信息采集')
body('自提订单需要补充以下取货信息：')
for x in ['取货车辆车牌号','司机姓名','联系电话','计划取货时间','联系人姓名','联系人联系方式']:
    bullet(x)
body('实际取货后，还应保留实际取货时间，计划时间与实际时间分开记录。')

h1('九 系统联动要求')
table(['系统','主要职责'],[
 ['OMS','订单接单、时间确认、履约计划更新'],
 ['OA','延期申请、争议时效协调及升级审批'],
 ['企业微信','待确认、待审批及处理结果提醒'],
 ['库存系统','库存量、占用量和可使用量查询'],
 ['物流平台','已对接平台自动同步物流信息；专车和自提订单人工补充资料'],
],[3.0,12.4])
body('所有时间字段应区分“客户要求时间”“预计时间”和“实际完成时间”。未产生实际业务事件时，不得提前展示为已完成。')

h1('十 待确认事项')
for x in [
 '快速接单可配置的监控节点、适用客户和启用条件。',
 '签约主体对应的费用项目、收费标准、费用承担方及结算方式。',
 'OA 审批接口、审批人替补机制及审批超时处理方式。',
 '企业微信通知对象、触发节点和重复提醒规则。',
 '库存充足的判定口径，以及库存接口异常时的接单处理方式。',
 '重点客户名单维护方式及汇总指标口径。'
]: bullet(x)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run('订单履约时效调整会议纪要'), size=9, color='666666')
footer.add_run('  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)

doc.core_properties.title = '订单履约时效调整会议纪要'
doc.core_properties.subject = '客户分级、接单时效、库存校验与物流资料采集'
doc.core_properties.author = '业务会议纪要'
doc.save(OUT)
print(OUT)
