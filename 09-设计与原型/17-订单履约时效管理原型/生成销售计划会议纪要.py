from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).with_name('销售计划管理调整会议纪要.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
sec.left_margin, sec.right_margin = Cm(2.4), Cm(2.4)

def font(run, size=None, bold=None, color='000000'):
    run.font.name = '微软雅黑'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

normal = doc.styles['Normal']
normal.font.name = '微软雅黑'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(0,0,0)
normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.28

for name,size,before,after in [('Title',22,0,18),('Heading 1',15,16,8),('Heading 2',12,10,5)]:
    s=doc.styles[name]; s.font.name='微软雅黑'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
    s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
    pPr=s._element.get_or_add_pPr(); border=pPr.find(qn('w:pBdr'))
    if border is not None: pPr.remove(border)

if '正文要点' not in doc.styles:
    s=doc.styles.add_style('正文要点',WD_STYLE_TYPE.PARAGRAPH); s.base_style=normal
    s.paragraph_format.left_indent=Cm(.7); s.paragraph_format.first_line_indent=Cm(-.35); s.paragraph_format.space_after=Pt(2.5)

def title(text):
    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pPr=p._p.get_or_add_pPr(); border=pPr.find(qn('w:pBdr'))
    if border is not None: pPr.remove(border)
    font(p.add_run(text),22,True)

def para(text, lead=None):
    p=doc.add_paragraph()
    if lead and text.startswith(lead): font(p.add_run(lead),bold=True); font(p.add_run(text[len(lead):]))
    else: font(p.add_run(text))
    return p

def intro(text):
    p=para(text); p.paragraph_format.line_spacing=1.5; p.paragraph_format.space_after=Pt(12)

def h1(text): doc.add_paragraph(text,style='Heading 1')
def h2(text): doc.add_paragraph(text,style='Heading 2')
def bullet(text,level=0):
    p=doc.add_paragraph(style='正文要点'); p.paragraph_format.left_indent=Cm(.7+level*.55); p.paragraph_format.first_line_indent=Cm(-.35)
    font(p.add_run('• '+text)); return p
def number(n,text):
    p=doc.add_paragraph(style='正文要点'); p.paragraph_format.left_indent=Cm(.8); p.paragraph_format.first_line_indent=Cm(-.55)
    font(p.add_run(f'{n}. '),bold=True); font(p.add_run(text)); return p

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); e=pr.find(qn('w:shd'))
    if e is None: e=OxmlElement('w:shd'); pr.append(e)
    e.set(qn('w:fill'),fill)
def cell_margin(cell):
    pr=cell._tc.get_or_add_tcPr(); m=pr.first_child_found_in('w:tcMar')
    if m is None: m=OxmlElement('w:tcMar'); pr.append(m)
    for k,v in [('top',105),('start',115),('bottom',105),('end',115)]:
        e=m.find(qn('w:'+k))
        if e is None: e=OxmlElement('w:'+k); m.append(e)
        e.set(qn('w:w'),str(v)); e.set(qn('w:type'),'dxa')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,x in enumerate(headers):
        c=t.rows[0].cells[i]; c.width=Cm(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade(c,'1F4E78'); cell_margin(c)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(x),9.5,True,'FFFFFF')
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,x in enumerate(row):
            c=cells[i]; c.width=Cm(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_margin(c)
            if ri%2: shade(c,'F3F7FA')
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i<2 else WD_ALIGN_PARAGRAPH.LEFT; font(p.add_run(str(x)),9.2)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

title('销售计划管理调整会议纪要')
intro('本次会议围绕销售计划的填报周期、线上线下管理口径、审核流程、销售数据参考、加急调整及达成分析进行梳理。系统需要形成从商品维护、计划填报、分级审核、计划调整到订单关联和达成分析的完整闭环，并保留全过程记录。')

h1('一 管理目标与计划范围')
bullet('同时支持线上销售计划和线下销售计划。')
bullet('计划分为月度计划和周计划两个管理维度。')
bullet('计划数据需要与实际销售订单关联，用于计算计划兑现情况和达成率。')
bullet('只有在“可填报商品”范围内维护的商品，才允许进入销售计划填报。')

h1('二 填报周期与时间控制')
table(['计划类型','填报节点','系统控制'],[
 ['月度计划','每月20日进行填报','开放填报前提前展示预警；超过截止时间后不允许继续填报'],
 ['周计划','每周五前完成填报','临近截止时间展示提醒；超过截止时间后不允许继续填报'],
],[3.0,4.1,8.3])
bullet('系统应清晰展示填报状态，包括未开始、可填报、即将截止、已截止、审核中、已确认。')
bullet('填报截止后如需补报或修改，应通过授权或调整流程处理，不直接覆盖原计划。')
bullet('月度计划具体开放日期、截止时点，以及周计划“周五前”的精确截止时间需要进一步确认。')

h1('三 月度计划填报与审核')
h2('线下月度计划')
bullet('以“省区、商品”为最小填报单位。')
bullet('省区提交计划后，逐级汇总至大区；大区汇总后提交中心。')
bullet('计划按组织层级逐级审核，最终由严总确认。')
para('建议系统同时保留省区原始填报值、大区汇总值、中心汇总值及各级审核意见，便于追踪汇总来源。')
h2('线上月度计划')
bullet('按照线上渠道进行计划填报和汇总。')
bullet('线上月度计划最终由丹总确认。')

h1('四 周计划填报与审核')
h2('线下周计划')
bullet('以“客户、商品”为最小填报单位。')
bullet('按照“省区、客户、商品”进行填报。')
bullet('省区提交后汇总至大区，大区汇总后提交中心。')
bullet('计划按组织层级逐级审核，最终由严总确认。')
h2('线上周计划')
bullet('按照线上渠道进行计划填报和汇总。')
bullet('线上周计划最终由丹总确认。')

doc.add_page_break()
h1('五 线上与线下销售数据参考')
para('填报计划时，系统应展示历史销售情况作为参考，但参考数据不直接替代填报结果。')
table(['业务类型','数据来源','参考维度'],[
 ['线上计划','旺店通','按线上渠道展示历史销售情况'],
 ['线下计划','OMS订单','支持查看近3个月、近2个月、近1个月及自定义时间区间的销售情况'],
],[3.1,3.5,8.8])
bullet('参考数据应明确统计时间范围、订单口径及数据更新时间。')
bullet('系统应支持从计划填报页面直接查看相关销售趋势和明细，减少跨系统查询。')
bullet('历史销售数据缺失或尚未同步时，应显示“待同步”或“暂无数据”，不得以零销量代替。')

h1('六 计划调整与加急审批')
para('产品变化或外部原因导致原计划无法继续执行时，允许快速调整计划，但必须保留原计划、调整内容、原因、审批过程和最终结果。')
h2('调整类型')
bullet('加急增加数量。')
bullet('加急减少数量。')
h2('调整流程')
for i,x in enumerate([
 '发起人选择原始计划及需要调整的计划明细。',
 '填写增加或减少的数量、调整原因及必要的附件。',
 '系统记录调整前数量、调整后数量、差异数量和发起时间。',
 '提交 OA 审批：线上计划由丹总审批，线下计划由严总审批。',
 '审批通过后生成新的计划版本，并按调整后的计划继续执行。',
 '审批未通过时保留原计划不变，同时记录驳回原因。'
],1): number(i,x)
para('任何调整都不得覆盖历史数据。计划页面应能够查看版本变化及每次调整记录。')

h1('七 计划完成情况与达成分析')
bullet('已完成计划分别按照月度计划和周计划进行归档和关联。')
bullet('计划与实际销售订单建立对应关系，支持从计划查看关联订单，也支持从订单查看所属计划。')
bullet('系统对比计划数量、实际完成数量和差异数量，并展示达成率或兑现率。')
bullet('支持按省区、客户、商品等多个维度查看计划与实际达成情况。')
bullet('线上计划还应支持按渠道查看；线下计划应支持大区、中心层级汇总。')
bullet('分析页面应支持从汇总指标下钻至计划明细和销售订单明细。')

h1('八 可填报商品维护')
bullet('建立独立的“可填报商品”维护功能。')
bullet('只有已启用的商品允许在月度计划和周计划中填报。')
bullet('商品维护应记录商品编码、商品名称、适用渠道、启用状态、生效时间和失效时间。')
bullet('商品停用后不影响历史计划查看，但不得继续用于新计划填报。')
bullet('商品范围发生调整时，应保留变更记录。')

h1('九 计划与销售订单关联')
bullet('月计划和周计划均需关联对应的销售订单。')
bullet('订单关联应根据计划周期、线上或线下类型、组织或渠道、客户及商品等条件进行匹配。')
bullet('系统应支持自动匹配，并允许有权限的人员处理未匹配或匹配异常的订单。')
bullet('同一订单或订单明细是否允许同时计入月计划和周计划，需要形成明确口径。')
bullet('取消、退款、退货、换货等订单变化是否回冲完成数量，需要纳入兑现率计算规则。')

h1('十 核心业务规则汇总')
table(['管理事项','线下','线上'],[
 ['月度最小单位','省区、商品','渠道'],
 ['周计划最小单位','客户、商品；按省区上报','渠道'],
 ['最终确认人','严总','丹总'],
 ['销售参考来源','OMS订单','旺店通'],
 ['加急调整审批','严总','丹总'],
 ['分析维度','中心、大区、省区、客户、商品','渠道、商品'],
],[4.0,5.7,5.7])

h1('十一 待确认事项')
for x in [
 '月度计划每月20日对应的计划月份，以及填报开放时间和准确截止时点。',
 '周计划对应的自然周范围，以及周五的准确截止时点。',
 '省区、大区、中心各级审核人的角色配置和退回规则。',
 '计划达成率、兑现率的计算公式、统计数量口径和金额口径。',
 '月计划与周计划同时关联同一订单时的统计规则。',
 '取消、退款、退货、换货及跨期订单对实际完成量的影响。',
 '加急调整是否允许多次发起，以及审批期间是否锁定原计划。',
 '旺店通和 OMS 数据同步频率、异常补偿及数据更新时间展示规则。'
]: bullet(x)

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run('销售计划管理调整会议纪要'),9,color='666666'); footer.add_run('  |  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title='销售计划管理调整会议纪要'
doc.core_properties.subject='月度和周销售计划填报 审核 调整与达成分析'
doc.core_properties.author='业务会议纪要'
doc.save(OUT)
print(OUT)
