from pathlib import Path
import openpyxl,zipfile
root=Path(__file__).parent
for name,plan in [('导出验收-已发货.xlsx','PL-EXPORT-SHIP'),('导出验收-搜索结果.xlsx','PL-DEMO-001')]:
    path=root/name
    with zipfile.ZipFile(path) as z: assert z.testzip() is None
    w=openpyxl.load_workbook(path,data_only=False)
    assert w.sheetnames==['发货计划','商品明细']
    s=w['发货计划'];assert s.max_row==2;assert s['A2'].value==plan;assert s.freeze_panes=='A2'
    assert s.auto_filter.ref=='A1:AG2'
    assert w['商品明细'].max_row==3
    if plan=='PL-EXPORT-SHIP':
        assert s['D2'].value=='已发货'
        assert s['E2'].value.startswith('=HYPERLINK') and s['E2'].data_type=='s'
        assert s['Y2'].value=='00001234567' and s['Y2'].data_type=='s'
        assert s['I2'].value==40 and s['I2'].data_type=='n'
        assert '\n' in s['AE2'].value and '& < >' in s['AE2'].value
    assert not any(c.data_type=='f' for sheet in w for row in sheet for c in row)
print('PASS: valid XLSX/CRC; filtered plan and detail rows; frozen headers and auto-filter; numeric quantities; leading-zero tracking IDs; multiline text; user strings cannot become formulas.')
