from pathlib import Path
import json,zipfile,xml.etree.ElementTree as ET
import openpyxl
root=Path('E:/cxy/99-附件与参考');out=Path('E:/cxy/outputs/logistics-review');out.mkdir(parents=True,exist_ok=True)
data={}
for name in ['跨越物流.xlsx','顺丰常温.xlsx','顺丰冷运.xlsx','中通冷链.xlsx']:
    w=openpyxl.load_workbook(root/name,data_only=False)
    sheets=[]
    for s in w:
        rows=[]
        for row in s:
            vals=[{'cell':c.coordinate,'value':c.value} for c in row if c.value is not None]
            if vals: rows.append(vals)
        entry={'sheet':s.title,'size':[s.max_row,s.max_column],'merged':[str(x) for x in s.merged_cells.ranges],'rows':rows,'images':len(s._images)}
        sheets.append(entry)
        print(name,s.title,entry['size'],'nonempty rows',len(rows),'images',entry['images'])
    data[name]=sheets
    with zipfile.ZipFile(root/name) as z:
        for m in z.namelist():
            if m.startswith('xl/media/'):
                p=out/(Path(name).stem+'-'+Path(m).name);p.write_bytes(z.read(m))
with zipfile.ZipFile(root/'物流计费明细.docx') as z:
    x=ET.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    paras=[''.join(p.itertext()) for p in []]
    paras=[''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in x.findall('.//w:p',ns)]
    data['物流计费明细.docx']=paras
    print('DOCX',len(paras),'paragraphs')
    for m in z.namelist():
        if m.startswith('word/media/'):(out/('计费明细-'+Path(m).name)).write_bytes(z.read(m))
(out/'extracted.json').write_text(json.dumps(data,ensure_ascii=False,default=str,indent=2),encoding='utf-8')
(out/'doc-text.txt').write_text('\n'.join(paras),encoding='utf-8')
print('DOC TEXT:\n'+'\n'.join(paras))
