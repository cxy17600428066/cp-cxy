from pathlib import Path
import sys
from html import escape
from docx import Document
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader
from pdf2image import convert_from_path
from PIL import Image,ImageOps,ImageDraw

base=Path(sys.argv[2] if len(sys.argv)>2 else 'E:/cxy/outputs/plan-prd/render');base.mkdir(parents=True,exist_ok=True)
src=Path(sys.argv[1] if len(sys.argv)>1 else 'E:/cxy/04-需求管理/PRD-生产管理计划管理-开发交付版-V1.0-20261027.docx')
d=Document(src)
font='PRDChinese'
pdfmetrics.registerFont(TTFont(font,'C:/Windows/Fonts/msyh.ttc',subfontIndex=0))
styles={k:ParagraphStyle(k,fontName=font,fontSize=size,leading=leading,spaceAfter=6,wordWrap='CJK',keepWithNext=keep) for k,size,leading,keep in [('Normal',11,14,False),('Title',24,31,True),('Subtitle',12,17,True),('Heading 1',17,23,True),('Heading 2',12,17,True),('Cell',9.5,12,False)]}
story=[];pars={p._p:p for p in d.paragraphs};tables={t._tbl:t for t in d.tables}
for e in d.element.body:
 if e.tag==qn('w:p'):
  if e.xpath('.//w:br[@w:type="page"]'):story.append(PageBreak());continue
  p=pars[e]
  if not p.text.strip():story.append(Spacer(1,8));continue
  style=styles.get(p.style.name,styles['Normal'])
  story.append(Paragraph(escape(p.text),style))
 elif e.tag==qn('w:tbl'):
  t=tables[e];rows=[[Paragraph(escape(c.text).replace('\n','<br/>'),styles['Cell']) for c in r.cells] for r in t.rows]
  widths=[c.width.inches*72 for c in t.columns]
  total=sum(widths);widths=[w*504/total for w in widths]
  table=Table(rows,colWidths=widths,repeatRows=1,hAlign='CENTER')
  table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.HexColor('#D9D9D9')),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E8EFF7')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F7F9FC')]),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
  story.append(table);story.append(Spacer(1,6))
def foot(c,doc):
 c.setFont(font,8);c.drawRightString(558,26,'排版核对  '+str(doc.page))
pdf=base/'layout-check.pdf';SimpleDocTemplate(str(pdf),pagesize=(612,792),leftMargin=54,rightMargin=54,topMargin=50,bottomMargin=50).build(story,onFirstPage=foot,onLaterPages=foot)
reader=PdfReader(str(pdf));print('Alternative layout pages',len(reader.pages))
for i,p in enumerate(reader.pages,1):
 text=p.extract_text();print(i,len(text),text[:55].replace('\n',' '))
images=convert_from_path(str(pdf),dpi=100,output_folder=str(base),fmt='png',output_file='page',paths_only=True)
for start in range(0,len(images),4):
 thumbs=[]
 for idx,path in enumerate(images[start:start+4],start+1):
  im=Image.open(path).convert('RGB');im.thumbnail((408,528));thumb=Image.new('RGB',(428,558),'#d7dce4');thumb.paste(im,((428-im.width)//2,20));ImageDraw.Draw(thumb).text((12,540),'PAGE '+str(idx),fill='black');thumbs.append(thumb)
 sheet=Image.new('RGB',(856,1116),'white')
 for j,im in enumerate(thumbs):sheet.paste(im,((j%2)*428,(j//2)*558))
 sheet.save(base/('contact-'+str(start//4+1)+'.png'))
print('Word structure',len(d.paragraphs),'paragraphs',len(d.tables),'tables')
assert len(d.tables)>=11
assert all('AC'+str(i).zfill(2) in '\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells) for i in range(1,15))
print('PASS offline content; native Word pagination not verified')
