from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
src=Path('E:/cxy/04-需求管理/PRD-生产管理计划管理-开发交付版-V1.0-20261027.docx')
out=src.with_name('PRD-生产管理计划管理-开发交付版-V1.1-20261027.docx')
d=Document(src)
def change(par,old,new):
 if old in par.text:
  text=par.text.replace(old,new)
  if par.runs:
   par.runs[0].text=text
   for r in par.runs[1:]:r.text=''
  else:par.add_run(text)
for p in d.paragraphs:
 change(p,'V1.0','V1.1')
 change(p,'核心验收AC01至AC24通过','核心验收AC01至AC30通过')
 change(p,'当前执行版本；内容或版本冲突','当前执行版本；审批中的数量修订允许由本级审核人按第21章保存，但不代表确认生效。内容或版本冲突')
 change(p,'审核、确认及调整生效使用事务和乐观锁。','审核、审批数量修订、确认及调整生效使用事务和乐观锁。')
 if p.text.startswith('通过及驳回均需填写意见。'):
  p.add_run(' 本级审核人员可先修订最小客户商品或渠道商品明细的数量，填写原因并保存，仍待本级审核；按第21章保留原提交及修订工作版本。')
for t in d.tables:
 for row in t.rows:
  cells=row.cells
  if cells[0].text in ['大区审核','中心审核','严总','线上领导','丹总'] and len(cells)==3 and t.rows[0].cells[1].text=='可操作事项':
   cells[1].text+='；本级待审明细数量修订'
  if cells[0].text=='待本级销售审核' and len(cells)==3:
   cells[1].text='查看、本级待审数量修订及审核；首级未审时整批撤回'
  for c in cells:
   for p in c.paragraphs:change(p,'AC01至AC24','AC01至AC30')
for section in d.sections:
 for p in section.footer.paragraphs:change(p,'V1.0','V1.1')
def p(text):d.add_paragraph(text)
def h(text):d.add_heading(text,2)
def table(headers,rows):
 t=d.add_table(rows=1,cols=3);t.autofit=False
 for col,w in zip(t.columns,[1.25,2.6,3.15]):col.width=Inches(w)
 for c,text in zip(t.rows[0].cells,headers):c.text=text
 for vals in rows:
  for c,text in zip(t.add_row().cells,vals):c.text=text
 sample=d.tables[0]
 for i,r in enumerate(t.rows):
  if i==0:
   x=OxmlElement('w:tblHeader');r._tr.get_or_add_trPr().append(x)
  for c in r.cells:
   c._tc.get_or_add_tcPr().append(deepcopy(sample.rows[min(i,1)].cells[0]._tc.get_or_add_tcPr())) if False else None
   border=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');border.append(x)
   c._tc.get_or_add_tcPr().append(border)
   if i==0:
    x=OxmlElement('w:shd');x.set(qn('w:fill'),'E8EFF7');c._tc.get_or_add_tcPr().append(x)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(4)
    for run in par.runs:run.font.size=Pt(9.5);run.bold=i==0
 d.add_paragraph()
d.add_page_break();d.add_heading('21 审批节点商品数量修订',1)
p('V1.1新增。大区审核、中心审核、严总、线上领导、丹总可以在自己当前待审的完整范围内，修改最小客户商品明细的数量；线上对应渠道商品明细。修订原因必填，保存后仍待本级审核，不自动通过，不提前生效。计划部门仍只有任务确认权，不获得本功能的销售修订权。')
h('权限与操作边界 P0')
p('本次开放数量修改，不允许换客户、商品、周期、组织或单位。只允许本级待审、未确认生效、未归档且未锁定的原计划。已生效计划及独立变更申请继续走原有调整流程。同一完整范围含未完成提交、驳回、待处理变更或版本冲突时，不放开修订。')
p('从审核汇总的“修改客户商品数量”或“修改渠道商品数量”进入，列表展示销售原填报量、当前待审量和版本；逐条修改数量与原因后保存。保存刷新完整汇总，审核动作仍须单独执行。对整份完整范围再次执行权限和版本校验，不允许保存时遗漏或覆盖其他人的更新。')
h('原始快照与审核工作版本')
p('原提交批次及其明细快照保持不可变。审批修订生成独立事件和工作快照，增加明细版本，但不改变当前审核节点。后续审核完整性及计划部门确认均核验工作快照，而不是篡改原提交快照来绕过校验。')
p('保存修订事件至少包括：修订id、原计划和明细id、原提交批次、当前审批节点、真实登录人员id、时间、修订前后数量、前后版本、完整前后内容及原因。每次领导审核另存该节点实际审核的数量与版本快照。下级已审记录只证明当时的版本，不自动补造对上级新数量的批准；流程从当前节点继续。')
h('同步与最终生效')
p('销售、各审批节点、计划部门的列表和详情使用同一最新工作版本，并可查看全部审批修订原因及历史。已审核列表同时区分当时审核数量和当前待审数量。原生效需求、销售达成及执行版本仍以计划部门最终确认结果为准。')
p('正式系统采用统一服务端持久化，写请求携带原明细版本、完整范围签名和幂等键。变更事件或轮询通知其他终端刷新，不能静默覆盖未保存输入。原型通过共享数据及同源窗口更新提示演示同步，不代表已经部署实时消息服务。')
d.add_page_break();d.add_heading('22 修订接口及补充验收',1)
table(['类别','契约或场景','要求或预期'],[
('写请求建议','明细id、修订后数量、原因、预期版本、完整范围签名、幂等键','操作者及节点取服务端登录态；保存修订不调用审核通过'),
('成功响应建议','新明细版本、工作快照版本、修订id、范围汇总、当前待审节点','各节点按相同版本刷新；保留原提交快照'),
('失败响应建议','无权限、状态变化、范围不完整、版本冲突、非法数量、原因缺失、保存失败','不修改任何数量、快照或日志，保留用户输入'),
('AC25','五种审批角色分别修订本级待审商品数量','均允许合法数量修订；保存后节点不变，仍未生效'),
('AC26','保存缺少原因、0、小数或未变化数量','均阻断；原数据及提交快照保持不变'),
('AC27','大区、中心、严总依次修改同一明细','每次事件独立留痕；下级能查当前量和原因，上级审核修订后量'),
('AC28','线上领导及丹总依次修改渠道商品明细','线上维度和权限不变；计划部门确认最终修订量'),
('AC29','越级、跨业务、已生效、锁定或过期版本修改','全部阻断；不能利用修改按钮越权或覆盖新版本'),
('AC30','保存失败、重复请求及原始审批快照核对','保存失败完整回滚，重复不增版本；原提交及既有审核数量未覆盖')])
h('实现注意事项')
p('审批修订事件与“审核通过”事件必须区分类型，不能把一次保存误记为本级已审核。历史审批查询不能要求所有下级审批版本都等于最终工作版本，否则会错误显示缺失记录。应按提交批次和节点快照还原各级真实审核情况。')
p('本次不新增销售重提、不跳过后续销售节点、不赋予计划部门销售改量权限。后续若需开放客户商品更换、跨周期变更或已生效计划由领导直接修改，应另行确认业务流程，不能沿用本次数量修订接口。')
d.save(out);print(out);print('tables',len(d.tables))
