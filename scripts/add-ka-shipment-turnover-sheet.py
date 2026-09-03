# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


SOURCE_PATH = Path(r"E:\cxy\07-数据与指标\outputs\KA销售新增统计字段模板-智能看板.xlsx")
OUTPUT_PATH = Path(r"E:\cxy\07-数据与指标\outputs\KA销售新增统计字段模板-大盘出货周转.xlsx")
SHEET_NAME = "KA大盘出货&周转"

COLORS = {
    "ink": "17202A",
    "muted": "5B6675",
    "line": "D8DEE7",
    "bg": "F6F8FB",
    "panel": "FFFFFF",
    "header": "D9EAF7",
    "section": "EAF0F8",
    "blue": "1D4ED8",
    "green": "15803D",
    "amber": "B45309",
    "red": "B91C1C",
}

thin = Side(style="thin", color=COLORS["line"])
border = Border(left=thin, right=thin, top=thin, bottom=thin)

SALES = "'销售明细'"
INV = "'库存明细'"
SHIP = "'出货退货'"


def style_range(ws, cell_range: str, fill: str = "FFFFFF") -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = Font(name="微软雅黑", size=10, color=COLORS["ink"])


def section(ws, cell_range: str, title: str) -> None:
    ws.merge_cells(cell_range)
    style_range(ws, cell_range, COLORS["section"])
    cell = ws[cell_range.split(":")[0]]
    cell.value = title
    cell.font = Font(name="微软雅黑", bold=True, size=12, color=COLORS["ink"])
    cell.alignment = Alignment(horizontal="left", vertical="center")


def header(ws, row: int, start_col: int, titles: list[str]) -> None:
    for i, title in enumerate(titles, start_col):
        cell = ws.cell(row=row, column=i)
        cell.value = title
        cell.fill = PatternFill("solid", fgColor=COLORS["header"])
        cell.font = Font(name="微软雅黑", bold=True, size=10, color=COLORS["ink"])
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_validation(ws, cell: str, values: list[str]) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[cell])


def opt_condition(sheet: str, col: str, param: str) -> str:
    return f'--((({param}="全部")+({sheet}!${col}$2:${col}$500={param}))>0)'


def sales_cond(date_ref="$B$3") -> str:
    return (
        f"({SALES}!$A$2:$A$500={date_ref})*"
        f"{opt_condition(SALES, 'B', '$D$3')}*"
        f"{opt_condition(SALES, 'C', '$F$3')}*"
        f"{opt_condition(SALES, 'D', '$H$3')}"
    )


def ship_cond(date_ref="$B$3") -> str:
    return (
        f"({SHIP}!$A$2:$A$500={date_ref})*"
        f"{opt_condition(SHIP, 'B', '$D$3')}*"
        f"{opt_condition(SHIP, 'C', '$F$3')}*"
        f"{opt_condition(SHIP, 'D', '$H$3')}"
    )


def inv_cond(date_ref="$B$3") -> str:
    return (
        f"({INV}!$A$2:$A$500={date_ref})*"
        f"{opt_condition(INV, 'B', '$D$3')}*"
        f"{opt_condition(INV, 'C', '$F$3')}*"
        f"{opt_condition(INV, 'D', '$H$3')}"
    )


def sumproduct(condition: str, sheet: str, col: str) -> str:
    return f"=SUMPRODUCT({condition}*{sheet}!${col}$2:${col}$500)"


def build() -> None:
    wb = load_workbook(SOURCE_PATH)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME, 1)
    ws.sheet_view.showGridLines = False

    for col in range(1, 25):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["E"].width = 24
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["K"].width = 24
    ws.column_dimensions["Q"].width = 24
    for row in range(1, 80):
        ws.row_dimensions[row].height = 22
    for row in ws["A1:X80"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=COLORS["bg"])

    ws.merge_cells("A1:X1")
    ws["A1"] = "KA大盘出货数据 & POS库存周转汇总"
    ws["A1"].font = Font(name="微软雅黑", bold=True, size=18, color=COLORS["ink"])
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    params = [
        ("B2", "统计日期", "B3", "=MAX('销售明细'!A2:A500)"),
        ("D2", "KA系统", "D3", "全部"),
        ("F2", "大区", "F3", "全部"),
        ("H2", "省份", "H3", "全部"),
    ]
    for label_cell, label, value_cell, value in params:
        ws[label_cell] = label
        ws[label_cell].font = Font(name="微软雅黑", bold=True, color=COLORS["muted"], size=9)
        ws[value_cell] = value
        ws[value_cell].fill = PatternFill("solid", fgColor=COLORS["panel"])
        ws[value_cell].font = Font(name="微软雅黑", bold=True, color=COLORS["ink"], size=10)
        ws[value_cell].alignment = Alignment(horizontal="center", vertical="center")
        ws[value_cell].border = border
    ws["B3"].number_format = "yyyy-mm-dd"
    add_validation(ws, "D3", ["全部", "大润发", "盒马", "华润", "沃尔玛", "永辉", "天虹", "比优特", "家得福"])
    add_validation(ws, "F3", ["全部", "华东", "华南", "华北", "华中", "西北", "东北"])
    add_validation(ws, "H3", ["全部", "上海", "江苏", "广东", "北京", "湖北", "浙江", "黑龙江", "天津"])

    # 1. KA大盘出货数据 by system / quantity / amount.
    section(ws, "B5:I5", "1、KA大盘出货数据：系统每天出货数量、出货金额（by系统/by数量/by金额）")
    header(ws, 6, 2, ["KA系统", "出货数量", "出货金额", "POS销售数量", "POS销售额", "出货/POS额", "退货数量", "退货金额"])
    systems = ["大润发", "盒马", "华润", "沃尔玛", "永辉", "天虹", "比优特", "家得福"]
    for i, system in enumerate(systems, 7):
        ws[f"B{i}"] = system
        sc = f"({SHIP}!$A$2:$A$500=$B$3)*--(({SHIP}!$B$2:$B$500=B{i})>0)*{opt_condition(SHIP, 'C', '$F$3')}*{opt_condition(SHIP, 'D', '$H$3')}"
        pc = f"({SALES}!$A$2:$A$500=$B$3)*--(({SALES}!$B$2:$B$500=B{i})>0)*{opt_condition(SALES, 'C', '$F$3')}*{opt_condition(SALES, 'D', '$H$3')}"
        ws[f"C{i}"] = sumproduct(sc, SHIP, "I")
        ws[f"D{i}"] = sumproduct(sc, SHIP, "J")
        ws[f"E{i}"] = sumproduct(pc, SALES, "N")
        ws[f"F{i}"] = sumproduct(pc, SALES, "O")
        ws[f"G{i}"] = f'=IFERROR(D{i}/F{i},0)'
        ws[f"H{i}"] = sumproduct(sc, SHIP, "K")
        ws[f"I{i}"] = sumproduct(sc, SHIP, "L")
    style_range(ws, "B7:I14", COLORS["panel"])
    for rng in ["C7:F14", "H7:I14"]:
        for row in ws[rng]:
            for cell in row:
                cell.number_format = "#,##0.0"
    for row in ws["G7:G14"]:
        row[0].number_format = "0.0%"

    # 1b. by item / channel.
    section(ws, "K5:R5", "1、KA大盘出货数据：by单品 / by渠道")
    header(ws, 6, 11, ["维度", "名称", "出货数量", "出货金额", "POS销售数量", "POS销售额", "出货/POS额", "备注"])
    items = [
        ("单品", "无骨鸡爪 柠檬味 158g", "4354168"),
        ("单品", "虎皮凤爪 白松露味", "4502138"),
        ("单品", "无骨鸡爪 柠檬味-沃尔玛", "25354863"),
        ("单品", "蒜香无骨鸡爪 230g", "431448153"),
        ("渠道", "POS", ""),
        ("渠道", "线下", ""),
        ("渠道", "京东到家", ""),
        ("渠道", "美团", ""),
        ("渠道", "天虹到家", ""),
    ]
    for i, (dim, name, sku) in enumerate(items, 7):
        ws[f"K{i}"] = dim
        ws[f"L{i}"] = name
        if dim == "单品":
            sc = f"{ship_cond()}*--(({SHIP}!$G$2:$G$500=\"{sku}\")>0)"
            pc = f"{sales_cond()}*--(({SALES}!$H$2:$H$500=\"{sku}\")>0)"
        else:
            sc = ship_cond()
            pc = f"{sales_cond()}*--(({SALES}!$M$2:$M$500=L{i})>0)"
        ws[f"M{i}"] = sumproduct(sc, SHIP, "I")
        ws[f"N{i}"] = sumproduct(sc, SHIP, "J")
        ws[f"O{i}"] = sumproduct(pc, SALES, "N")
        ws[f"P{i}"] = sumproduct(pc, SALES, "O")
        ws[f"Q{i}"] = f'=IFERROR(N{i}/P{i},0)'
        ws[f"R{i}"] = '=IF(K{0}="渠道","渠道出货需明细源支持","")'.format(i)
    style_range(ws, "K7:R15", COLORS["panel"])
    for rng in ["M7:P15"]:
        for row in ws[rng]:
            for cell in row:
                cell.number_format = "#,##0.0"
    for row in ws["Q7:Q15"]:
        row[0].number_format = "0.0%"

    # 2. POS comparison by item.
    section(ws, "B18:I18", "2、系统单品出货额和POS额的对比")
    header(ws, 19, 2, ["KA系统", "商品编码", "商品名称", "出货金额", "POS销售额", "金额差异", "出货/POS额", "状态"])
    sku_rows = [
        ("大润发", "4354168", "脱骨侠无骨鸡爪(柠檬味)"),
        ("大润发", "4502138", "脱骨侠虎皮凤爪(白松露味)"),
        ("沃尔玛", "25354863", "常温脱骨侠无骨鸡爪柠檬"),
        ("盒马", "431448153", "盒马工坊 蒜香无骨鸡爪 230g"),
        ("天虹", "v000110869", "脱骨侠双椒无骨鸭掌"),
        ("家得福", "10081319", "脱骨侠酸辣柠檬无骨鸡爪780g"),
    ]
    for i, (sys, sku, name) in enumerate(sku_rows, 20):
        ws[f"B{i}"] = sys
        ws[f"C{i}"] = sku
        ws[f"D{i}"] = name
        sc = f"{ship_cond()}*--(({SHIP}!$B$2:$B$500=B{i})>0)*--(({SHIP}!$G$2:$G$500=C{i})>0)"
        pc = f"{sales_cond()}*--(({SALES}!$B$2:$B$500=B{i})>0)*--(({SALES}!$H$2:$H$500=C{i})>0)"
        ws[f"E{i}"] = sumproduct(sc, SHIP, "J")
        ws[f"F{i}"] = sumproduct(pc, SALES, "O")
        ws[f"G{i}"] = f"=E{i}-F{i}"
        ws[f"H{i}"] = f'=IFERROR(E{i}/F{i},0)'
        ws[f"I{i}"] = f'=IF(E{i}=0,"缺出货源",IF(H{i}<0.8,"低于POS",IF(H{i}>1.2,"高于POS","正常")))'
    style_range(ws, "B20:I25", COLORS["panel"])
    for rng in ["E20:G25"]:
        for row in ws[rng]:
            for cell in row:
                cell.number_format = "#,##0.0"
    for row in ws["H20:H25"]:
        row[0].number_format = "0.0%"

    # 3. Regional SKU turnover.
    section(ws, "K18:R18", "POS数据：各大区区分各单品库存周转天数（数据日前7天平均）")
    header(ws, 19, 11, ["大区", "商品编码", "商品名称", "门店库存", "近7日平均销量", "周转天数", "超60天门店数", "风险"])
    turnover_rows = [
        ("华东", "4354168", "脱骨侠无骨鸡爪(柠檬味)"),
        ("华南", "4502138", "脱骨侠虎皮凤爪(白松露味)"),
        ("华南", "25354863", "常温脱骨侠无骨鸡爪柠檬"),
        ("华北", "5150673", "脱骨侠无骨鸡爪(柠檬味)158g"),
        ("华中", "431448153", "盒马工坊 蒜香无骨鸡爪 230g"),
        ("华东", "10081319", "脱骨侠酸辣柠檬无骨鸡爪780g"),
    ]
    for i, (region, sku, name) in enumerate(turnover_rows, 20):
        ws[f"K{i}"] = region
        ws[f"L{i}"] = sku
        ws[f"M{i}"] = name
        cond = f"{inv_cond()}*--(({INV}!$C$2:$C$500=K{i})>0)*--(({INV}!$H$2:$H$500=L{i})>0)"
        ws[f"N{i}"] = sumproduct(cond, INV, "J")
        ws[f"O{i}"] = sumproduct(cond, INV, "M")
        ws[f"P{i}"] = f'=IFERROR(N{i}/O{i},"")'
        ws[f"Q{i}"] = f'=COUNTIFS({INV}!$A$2:$A$500,$B$3,{INV}!$C$2:$C$500,K{i},{INV}!$H$2:$H$500,L{i},{INV}!$O$2:$O$500,"Y")'
        ws[f"R{i}"] = f'=IF(P{i}="","无动销",IF(P{i}>60,"高风险","正常"))'
    style_range(ws, "K20:R25", COLORS["panel"])
    for rng in ["N20:Q25"]:
        for row in ws[rng]:
            for cell in row:
                cell.number_format = "#,##0.0"
    ws.conditional_formatting.add("P20:P25", CellIsRule(operator="greaterThan", formula=["60"], fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(color=COLORS["red"], bold=True)))
    ws.conditional_formatting.add("P20:P25", IconSetRule("3TrafficLights1", "num", [30, 60, 120], showValue=True))

    # 4. Ranking stores >60 by region.
    section(ws, "B29:I29", "各大区库存周转天数超60天门店数排名")
    header(ws, 30, 2, ["排名", "大区", "超60天门店数", "门店库存", "近7日平均销量", "平均周转天数", "主要风险单品", "处理建议"])
    regions = ["华东", "华南", "华北", "华中", "东北"]
    for idx, region in enumerate(regions, 31):
        ws[f"K{idx}"] = region
        cond = f"{inv_cond()}*--(({INV}!$C$2:$C$500=K{idx})>0)"
        ws[f"L{idx}"] = f'=COUNTIFS({INV}!$A$2:$A$500,$B$3,{INV}!$C$2:$C$500,K{idx},{INV}!$O$2:$O$500,"Y")'
        ws[f"M{idx}"] = sumproduct(cond, INV, "J")
        ws[f"N{idx}"] = sumproduct(cond, INV, "M")
        ws[f"O{idx}"] = f'=IFERROR(M{idx}/N{idx},"")'
    for i in range(5):
        row = 31 + i
        ws[f"B{row}"] = i + 1
        ws[f"C{row}"] = f'=IFERROR(INDEX($K$31:$K$35,MATCH(LARGE($L$31:$L$35,B{row}),$L$31:$L$35,0)),"")'
        ws[f"D{row}"] = f'=IFERROR(INDEX($L$31:$L$35,MATCH(C{row},$K$31:$K$35,0)),0)'
        ws[f"E{row}"] = f'=IFERROR(INDEX($M$31:$M$35,MATCH(C{row},$K$31:$K$35,0)),0)'
        ws[f"F{row}"] = f'=IFERROR(INDEX($N$31:$N$35,MATCH(C{row},$K$31:$K$35,0)),0)'
        ws[f"G{row}"] = f'=IFERROR(INDEX($O$31:$O$35,MATCH(C{row},$K$31:$K$35,0)),"")'
        ws[f"H{row}"] = '=IFERROR(INDEX($M$20:$M$25,MATCH(MAXIFS($P$20:$P$25,$K$20:$K$25,C{0}),$P$20:$P$25,0)),"")'.format(row)
        ws[f"I{row}"] = f'=IF(D{row}>0,"优先清理无动销库存/调拨/促销","正常跟踪")'
    style_range(ws, "B31:I35", COLORS["panel"])
    for rng in ["D31:G35"]:
        for row in ws[rng]:
            for cell in row:
                cell.number_format = "#,##0.0"
    ws.conditional_formatting.add("D31:D35", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(color=COLORS["red"], bold=True)))

    # Notes.
    section(ws, "B39:I39", "口径说明")
    notes = [
        ["出货数量/金额", "取自 `出货退货` 页；当前样例中部分KA缺出货金额，展示为0或空，需补源。"],
        ["POS销售额", "取自 `销售明细` 页销售金额；沃尔玛未税/含税需通过含税口径字段区分。"],
        ["周转天数", "门店库存 / 近7日平均销量，近7日平均销量取自 `库存明细` 页，可由D-7至D-1销售推导。"],
        ["超60天门店数", "按统计日期、大区、商品编码，统计 `是否超60天=Y` 的门店数量。"],
    ]
    header(ws, 40, 2, ["项目", "说明"])
    for i, row in enumerate(notes, 41):
        ws[f"B{i}"] = row[0]
        ws[f"C{i}"] = row[1]
    style_range(ws, "B41:C44", COLORS["panel"])
    ws.column_dimensions["C"].width = 48

    # Charts.
    add_charts(ws)

    # Hide ranking helper columns K:O rows 31:35? Keep visible for audit but style muted.
    for row in ws["K31:O35"]:
        for cell in row:
            cell.font = Font(name="微软雅黑", size=8, color=COLORS["muted"])
            cell.fill = PatternFill("solid", fgColor=COLORS["bg"])

    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(OUTPUT_PATH)

    check = load_workbook(OUTPUT_PATH, data_only=False)
    assert SHEET_NAME in check.sheetnames
    ws2 = check[SHEET_NAME]
    assert ws2["B5"].value.startswith("1、KA大盘出货数据")
    assert len(ws2._charts) >= 3
    assert ws2["E20"].value.startswith("=SUMPRODUCT") or ws2["E20"].value == 0
    check.close()
    print(OUTPUT_PATH)


def add_charts(ws) -> None:
    chart1 = BarChart()
    chart1.title = "KA系统出货金额 vs POS销售额"
    chart1.y_axis.title = "金额"
    chart1.height = 7
    chart1.width = 15
    chart1.add_data(Reference(ws, min_col=4, min_row=6, max_col=6, max_row=14), titles_from_data=True)
    chart1.set_categories(Reference(ws, min_col=2, min_row=7, max_row=14))
    ws.add_chart(chart1, "B47")

    chart2 = BarChart()
    chart2.title = "单品出货/POS金额比"
    chart2.y_axis.title = "比例"
    chart2.height = 7
    chart2.width = 15
    chart2.add_data(Reference(ws, min_col=8, min_row=19, max_row=25), titles_from_data=True)
    chart2.set_categories(Reference(ws, min_col=4, min_row=20, max_row=25))
    ws.add_chart(chart2, "K29")

    chart3 = BarChart()
    chart3.type = "bar"
    chart3.title = "区域超60天门店数排名"
    chart3.y_axis.title = "大区"
    chart3.x_axis.title = "门店数"
    chart3.height = 7
    chart3.width = 15
    chart3.add_data(Reference(ws, min_col=4, min_row=30, max_row=35), titles_from_data=True)
    chart3.set_categories(Reference(ws, min_col=3, min_row=31, max_row=35))
    ws.add_chart(chart3, "K47")


if __name__ == "__main__":
    build()
