# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


SOURCE_PATH = Path(r"E:\cxy\07-数据与指标\outputs\KA销售新增统计字段模板-含日报看板.xlsx")
OUTPUT_PATH = Path(r"E:\cxy\07-数据与指标\outputs\KA销售新增统计字段模板-智能看板.xlsx")
SHEET_NAME = "智能日报看板"

COLORS = {
    "ink": "17202A",
    "muted": "5B6675",
    "line": "D8DEE7",
    "bg": "F6F8FB",
    "panel": "FFFFFF",
    "header": "EAF0F8",
    "section": "F0F4F8",
    "blue": "1D4ED8",
    "green": "15803D",
    "amber": "B45309",
    "red": "B91C1C",
    "purple": "6D28D9",
    "teal": "0891B2",
}

thin = Side(style="thin", color=COLORS["line"])
border = Border(left=thin, right=thin, top=thin, bottom=thin)

SALES = "'销售明细'"
INV = "'库存明细'"
SHIP = "'出货退货'"
STORE = "'门店维表'"


def fill_range(ws, cell_range: str, color: str) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=color)


def style_box(ws, cell_range: str, fill: str = "FFFFFF") -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def section(ws, cell_range: str, title: str) -> None:
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = title
    cell.font = Font(name="微软雅黑", bold=True, color=COLORS["ink"], size=11)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    style_box(ws, cell_range, COLORS["section"])


def header_row(ws, row: int, start_col: int, headers: list[str]) -> None:
    for offset, title in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + offset)
        cell.value = title
        cell.fill = PatternFill("solid", fgColor=COLORS["header"])
        cell.font = Font(name="微软雅黑", bold=True, color=COLORS["ink"], size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border


def table_body(ws, cell_range: str) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=COLORS["panel"])
            cell.border = border
            cell.font = Font(name="微软雅黑", size=10, color=COLORS["ink"])
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_validation(ws, cell: str, values: list[str]) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[cell])


def opt_condition(sheet: str, col: str, param: str) -> str:
    return f'--((({param}="全部")+({sheet}!${col}$2:${col}$500={param}))>0)'


def sales_condition(date_ref: str) -> str:
    return (
        f"({SALES}!$A$2:$A$500={date_ref})*"
        f"{opt_condition(SALES, 'B', '$D$3')}*"
        f"{opt_condition(SALES, 'C', '$F$3')}*"
        f"{opt_condition(SALES, 'D', '$H$3')}*"
        f"{opt_condition(SALES, 'L', '$J$3')}*"
        f"{opt_condition(SALES, 'M', '$L$3')}"
    )


def sales_sum(date_ref: str, value_col: str) -> str:
    return f"=SUMPRODUCT({sales_condition(date_ref)}*{SALES}!${value_col}$2:${value_col}$500)"


def inv_condition(date_ref: str) -> str:
    return (
        f"({INV}!$A$2:$A$500={date_ref})*"
        f"{opt_condition(INV, 'B', '$D$3')}*"
        f"{opt_condition(INV, 'C', '$F$3')}*"
        f"{opt_condition(INV, 'D', '$H$3')}"
    )


def inv_sum(date_ref: str, value_col: str) -> str:
    return f"=SUMPRODUCT({inv_condition(date_ref)}*{INV}!${value_col}$2:${value_col}$500)"


def ship_condition(date_ref: str) -> str:
    return (
        f"({SHIP}!$A$2:$A$500={date_ref})*"
        f"{opt_condition(SHIP, 'B', '$D$3')}*"
        f"{opt_condition(SHIP, 'C', '$F$3')}*"
        f"{opt_condition(SHIP, 'D', '$H$3')}"
    )


def ship_sum(date_ref: str, value_col: str) -> str:
    return f"=SUMPRODUCT({ship_condition(date_ref)}*{SHIP}!${value_col}$2:${value_col}$500)"


def kpi(ws, cell_range: str, title: str, formula: str, delta_formula: str, color: str) -> None:
    style_box(ws, cell_range, COLORS["panel"])
    top_left = ws[cell_range.split(":")[0]]
    r = top_left.row
    c = top_left.column
    ws.cell(r, c).value = title
    ws.cell(r, c).font = Font(name="微软雅黑", bold=True, color=COLORS["muted"], size=10)
    ws.cell(r + 1, c).value = formula
    ws.cell(r + 1, c).font = Font(name="微软雅黑", bold=True, color=color, size=18)
    ws.cell(r + 1, c).number_format = "#,##0"
    ws.cell(r + 3, c).value = delta_formula
    ws.cell(r + 3, c).font = Font(name="微软雅黑", bold=True, color=COLORS["muted"], size=9)
    for rr in range(r, r + 4):
        ws.cell(rr, c).alignment = Alignment(horizontal="left", vertical="center")


def delta_text(current_cell: str, previous_formula: str) -> str:
    return f'=IFERROR("环比昨日"&IF({current_cell}/{previous_formula}-1>=0,"↑","↓")&TEXT(ABS({current_cell}/{previous_formula}-1),"0.0%"),"环比昨日-")'


def build() -> None:
    wb = load_workbook(SOURCE_PATH)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME, 1)
    ws.sheet_view.showGridLines = False

    for col in range(1, 31):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 4
    ws.column_dimensions["H"].width = 15
    ws.column_dimensions["I"].width = 15
    ws.column_dimensions["J"].width = 15
    ws.column_dimensions["K"].width = 15
    ws.column_dimensions["L"].width = 15

    for row in range(1, 90):
        ws.row_dimensions[row].height = 22
    fill_range(ws, "A1:AD90", COLORS["bg"])

    ws.merge_cells("A1:L1")
    ws["A1"] = "KA智能日报看板"
    ws["A1"].font = Font(name="微软雅黑", bold=True, size=18, color=COLORS["ink"])
    ws["A2"] = "筛选参数"
    ws["A2"].font = Font(name="微软雅黑", bold=True, size=10, color=COLORS["muted"])

    params = [
        ("B2", "统计日期", "B3", "=MAX('销售明细'!A2:A500)"),
        ("D2", "KA系统", "D3", "全部"),
        ("F2", "大区", "F3", "全部"),
        ("H2", "省份", "H3", "全部"),
        ("J2", "温层", "J3", "全部"),
        ("L2", "渠道", "L3", "全部"),
    ]
    for label_cell, label, value_cell, value in params:
        ws[label_cell] = label
        ws[label_cell].font = Font(name="微软雅黑", bold=True, color=COLORS["muted"], size=9)
        ws[value_cell] = value
        ws[value_cell].font = Font(name="微软雅黑", bold=True, color=COLORS["ink"], size=10)
        ws[value_cell].fill = PatternFill("solid", fgColor=COLORS["panel"])
        ws[value_cell].border = border
        ws[value_cell].alignment = Alignment(horizontal="center", vertical="center")
    ws["B3"].number_format = "yyyy-mm-dd"
    add_validation(ws, "D3", ["全部", "大润发", "盒马", "华润", "沃尔玛", "永辉", "天虹", "比优特", "家得福"])
    add_validation(ws, "F3", ["全部", "华东", "华南", "华北", "华中", "西北", "东北"])
    add_validation(ws, "H3", ["全部", "上海", "江苏", "广东", "北京", "湖北", "浙江", "黑龙江", "天津"])
    add_validation(ws, "J3", ["全部", "常温", "冷藏"])
    add_validation(ws, "L3", ["全部", "POS", "线下", "线上", "京东到家", "美团", "天虹到家"])

    ws.merge_cells("A4:L4")
    ws["A4"] = "说明：销售、品项、渠道、门店TOP按销售明细汇总；库存与周转按库存明细汇总；出货/退货按出货退货页汇总。库存页暂不受温层、渠道筛选影响。"
    ws["A4"].font = Font(name="微软雅黑", size=9, color=COLORS["muted"])

    # KPI cards.
    kpi(ws, "B6:C9", "销售额", sales_sum("$B$3", "O"), delta_text("B7", sales_sum("$B$3-1", "O")[1:]), COLORS["blue"])
    kpi(ws, "D6:E9", "销量", sales_sum("$B$3", "N"), delta_text("D7", sales_sum("$B$3-1", "N")[1:]), COLORS["green"])
    kpi(ws, "F6:G9", "当月日均销量", f'=IFERROR(SUMPRODUCT((TEXT({SALES}!$A$2:$A$500,"yyyymm")=TEXT($B$3,"yyyymm"))*{opt_condition(SALES, "B", "$D$3")}*{opt_condition(SALES, "C", "$F$3")}*{opt_condition(SALES, "D", "$H$3")}*{opt_condition(SALES, "L", "$J$3")}*{opt_condition(SALES, "M", "$L$3")}*{SALES}!$N$2:$N$500)/DAY($B$3),0)', "", COLORS["purple"])
    kpi(ws, "H6:I9", "库存", inv_sum("$B$3", "J"), delta_text("H7", inv_sum("$B$3-1", "J")[1:]), COLORS["amber"])
    kpi(ws, "J6:K9", "有库存门店数", f'=IFERROR(ROWS(UNIQUE(FILTER({INV}!$E$2:$E$500,({inv_condition("$B$3")})*({INV}!$J$2:$J$500>0)))),0)', "", COLORS["teal"])
    kpi(ws, "L6:M9", "动销门店数", f'=IFERROR(ROWS(UNIQUE(FILTER({SALES}!$E$2:$E$500,({sales_condition("$B$3")})*({SALES}!$N$2:$N$500>0)))),0)', "", COLORS["red"])

    # Visible insight tables.
    section(ws, "B11:F11", "门店明细-动销TOP5")
    header_row(ws, 12, 2, ["门店名称", "销售数量", "销售金额", "库存数量", "周转天数"])
    for i in range(5):
        row = 13 + i
        helper_row = 44 + i
        ws[f"B{row}"] = f'=IFERROR(INDEX($Y$44:$Y$53,MATCH(LARGE($AA$44:$AA$53,ROWS($B$13:B{row})),$AA$44:$AA$53,0)),"")'
        ws[f"C{row}"] = f"=IFERROR(INDEX($Z$44:$Z$53,MATCH(B{row},$Y$44:$Y$53,0)),0)"
        ws[f"D{row}"] = f"=IFERROR(INDEX($AA$44:$AA$53,MATCH(B{row},$Y$44:$Y$53,0)),0)"
        ws[f"E{row}"] = f"=IFERROR(INDEX($AB$44:$AB$53,MATCH(B{row},$Y$44:$Y$53,0)),0)"
        ws[f"F{row}"] = f'=IFERROR(INDEX($AC$44:$AC$53,MATCH(B{row},$Y$44:$Y$53,0)),"")'
    table_body(ws, "B13:F17")

    section(ws, "H11:M11", "经营提醒")
    header_row(ws, 12, 8, ["提醒项", "结果", "建议动作"])
    alerts = [
        ("出货/POS", '=IFERROR(IF($J$24/$I$24<0.8,"偏低","正常"),"缺数据")', "低于80%时检查补货或出货源"),
        ("退货率", '=IFERROR(IF($L$24>0.05,"偏高","正常"),"缺数据")', "高于5%时按单品追踪原因"),
        ("周转风险", '=IF($K$7>0,"存在超60天门店","正常")', "优先处理库存大且无动销门店"),
        ("数据质量", '=IF(COUNTBLANK(\'销售明细\'!A2:O500)>200,"待补字段","可用")', "补齐日期、门店、商品、金额字段"),
    ]
    for i, row in enumerate(alerts, 13):
        ws[f"H{i}"], ws[f"I{i}"], ws[f"J{i}"] = row
    table_body(ws, "H13:J16")

    # Helper area used by charts and formulas.
    ws["O1"] = "图表/明细计算区"
    ws["O1"].font = Font(name="微软雅黑", bold=True, color=COLORS["muted"])
    build_helpers(ws)
    add_charts(ws)

    # Formatting.
    for cell_range in ["B7:M7", "C13:E17", "P4:R10", "P14:R22", "P26:R33", "P37:R42", "Y44:AC53"]:
        for row in ws[cell_range]:
            for cell in row:
                cell.number_format = "#,##0.0"
    for rng in ["B6:M9", "B11:F17", "H11:M16"]:
        style_box(ws, rng, COLORS["panel"])
    ws.conditional_formatting.add("F13:F17", CellIsRule(operator="greaterThan", formula=["60"], fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(color=COLORS["red"], bold=True)))
    ws.conditional_formatting.add("F13:F17", IconSetRule("3TrafficLights1", "num", [30, 60, 120], showValue=True))

    # Keep helper visible but off to the right, so users can audit formulas.
    for col in range(15, 30):
        ws.column_dimensions[get_column_letter(col)].width = 14

    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(OUTPUT_PATH)

    check = load_workbook(OUTPUT_PATH, data_only=False)
    assert SHEET_NAME in check.sheetnames
    assert len(check[SHEET_NAME]._charts) >= 5
    assert check[SHEET_NAME]["B7"].value.startswith("=SUMPRODUCT")
    check.close()
    print(OUTPUT_PATH)


def build_helpers(ws) -> None:
    # Trend helper.
    header_row(ws, 3, 16, ["日期", "销售额", "销量"])
    for i in range(7):
        row = 4 + i
        ws[f"P{row}"] = f"=$B$3-{6-i}"
        ws[f"P{row}"].number_format = "yyyy-mm-dd"
        ws[f"Q{row}"] = sales_sum(f"$P${row}", "O")
        ws[f"R{row}"] = sales_sum(f"$P${row}", "N")
    table_body(ws, "P4:R10")

    # System helper.
    header_row(ws, 13, 16, ["KA系统", "销售额", "销量"])
    systems = ["大润发", "盒马", "华润", "沃尔玛", "永辉", "天虹", "比优特", "家得福"]
    for i, system in enumerate(systems, 14):
        ws[f"P{i}"] = system
        cond = sales_condition("$B$3").replace('$D$3="全部")+(' + f"{SALES}!$B$2:$B$500=$D$3", f'TRUE)+({SALES}!$B$2:$B$500=P{i}')
        # System chart intentionally ignores the KA filter to show full contribution under other filters.
        cond = (
            f"({SALES}!$A$2:$A$500=$B$3)*"
            f"--(({SALES}!$B$2:$B$500=P{i})>0)*"
            f"{opt_condition(SALES, 'C', '$F$3')}*"
            f"{opt_condition(SALES, 'D', '$H$3')}*"
            f"{opt_condition(SALES, 'L', '$J$3')}*"
            f"{opt_condition(SALES, 'M', '$L$3')}"
        )
        ws[f"Q{i}"] = f"=SUMPRODUCT({cond}*{SALES}!$O$2:$O$500)"
        ws[f"R{i}"] = f"=SUMPRODUCT({cond}*{SALES}!$N$2:$N$500)"
    table_body(ws, "P14:R21")

    # Channel helper.
    header_row(ws, 25, 16, ["渠道", "销售额", "占比"])
    channels = ["POS", "线下", "线上", "京东到家", "美团", "天虹到家", "其他"]
    for i, channel in enumerate(channels, 26):
        ws[f"P{i}"] = channel
        if channel == "其他":
            channel_cond = f'--(ISNA(MATCH({SALES}!$M$2:$M$500,{{"POS","线下","线上","京东到家","美团","天虹到家"}},0)))'
        else:
            channel_cond = f'--(({SALES}!$M$2:$M$500=P{i})>0)'
        cond = (
            f"({SALES}!$A$2:$A$500=$B$3)*"
            f"{opt_condition(SALES, 'B', '$D$3')}*"
            f"{opt_condition(SALES, 'C', '$F$3')}*"
            f"{opt_condition(SALES, 'D', '$H$3')}*"
            f"{opt_condition(SALES, 'L', '$J$3')}*"
            f"{channel_cond}"
        )
        ws[f"Q{i}"] = f"=SUMPRODUCT({cond}*{SALES}!$O$2:$O$500)"
        ws[f"R{i}"] = f"=IFERROR(Q{i}/SUM($Q$26:$Q$32),0)"
    table_body(ws, "P26:R32")
    for row in range(26, 33):
        ws[f"R{row}"].number_format = "0.0%"

    # Shipment helper.
    header_row(ws, 23, 8, ["POS销售额", "出货金额", "退货金额", "退货率"])
    ws["H24"] = "汇总"
    ws["I24"] = sales_sum("$B$3", "O")
    ws["J24"] = ship_sum("$B$3", "J")
    ws["K24"] = ship_sum("$B$3", "L")
    ws["L24"] = "=IFERROR(K24/J24,0)"
    table_body(ws, "H24:L24")
    ws["L24"].number_format = "0.0%"

    # Inventory risk helper.
    header_row(ws, 36, 16, ["大区", "平均周转", "超60天门店"])
    regions = ["华东", "华南", "华北", "华中", "东北"]
    for i, region in enumerate(regions, 37):
        ws[f"P{i}"] = region
        cond = (
            f"({INV}!$A$2:$A$500=$B$3)*"
            f"--(({INV}!$C$2:$C$500=P{i})>0)*"
            f"{opt_condition(INV, 'B', '$D$3')}*"
            f"{opt_condition(INV, 'D', '$H$3')}"
        )
        ws[f"Q{i}"] = f"=IFERROR(SUMPRODUCT({cond}*{INV}!$N$2:$N$500)/SUMPRODUCT({cond}*--({INV}!$N$2:$N$500<>\"\")),0)"
        ws[f"R{i}"] = f"=COUNTIFS({INV}!$A$2:$A$500,$B$3,{INV}!$C$2:$C$500,P{i},{INV}!$O$2:$O$500,\"Y\")"
    table_body(ws, "P37:R41")

    # Store helper.
    header_row(ws, 43, 25, ["门店名称", "销售数量", "销售金额", "库存数量", "周转天数"])
    for i in range(10):
        row = 44 + i
        store_row = 2 + i
        ws[f"Y{row}"] = f"={STORE}!C{store_row}"
        ws[f"Z{row}"] = f"=SUMIFS({SALES}!$N$2:$N$500,{SALES}!$A$2:$A$500,$B$3,{SALES}!$F$2:$F$500,Y{row})"
        ws[f"AA{row}"] = f"=SUMIFS({SALES}!$O$2:$O$500,{SALES}!$A$2:$A$500,$B$3,{SALES}!$F$2:$F$500,Y{row})"
        ws[f"AB{row}"] = f"=SUMIFS({INV}!$J$2:$J$500,{INV}!$A$2:$A$500,$B$3,{INV}!$F$2:$F$500,Y{row})"
        ws[f"AC{row}"] = f'=IFERROR(AB{row}/Z{row},"")'
    table_body(ws, "Y44:AC53")

    # Product helper.
    header_row(ws, 45, 16, ["品项", "销售额", "销量"])
    for i in range(5):
        row = 46 + i
        sku_row = 2 + i
        ws[f"P{row}"] = f"='商品维表'!D{sku_row}"
        ws[f"Q{row}"] = f"=SUMIFS({SALES}!$O$2:$O$500,{SALES}!$A$2:$A$500,$B$3,{SALES}!$H$2:$H$500,'商品维表'!B{sku_row})"
        ws[f"R{row}"] = f"=SUMIFS({SALES}!$N$2:$N$500,{SALES}!$A$2:$A$500,$B$3,{SALES}!$H$2:$H$500,'商品维表'!B{sku_row})"
    table_body(ws, "P46:R50")


def add_charts(ws) -> None:
    trend = LineChart()
    trend.title = "近7日销售趋势"
    trend.y_axis.title = "金额 / 数量"
    trend.height = 7
    trend.width = 14
    trend.add_data(Reference(ws, min_col=17, min_row=3, max_col=18, max_row=10), titles_from_data=True)
    trend.set_categories(Reference(ws, min_col=16, min_row=4, max_row=10))
    ws.add_chart(trend, "B19")

    system = BarChart()
    system.title = "KA系统销售额贡献"
    system.y_axis.title = "销售额"
    system.height = 7
    system.width = 14
    system.add_data(Reference(ws, min_col=17, min_row=13, max_row=21), titles_from_data=True)
    system.set_categories(Reference(ws, min_col=16, min_row=14, max_row=21))
    ws.add_chart(system, "H19")

    channel = PieChart()
    channel.title = "渠道销售额占比"
    channel.height = 7
    channel.width = 14
    channel.add_data(Reference(ws, min_col=17, min_row=26, max_row=32), titles_from_data=False)
    channel.set_categories(Reference(ws, min_col=16, min_row=26, max_row=32))
    channel.dataLabels = DataLabelList()
    channel.dataLabels.showPercent = True
    ws.add_chart(channel, "B35")

    risk = BarChart()
    risk.type = "bar"
    risk.title = "区域周转风险"
    risk.height = 7
    risk.width = 14
    risk.add_data(Reference(ws, min_col=17, min_row=36, max_col=18, max_row=41), titles_from_data=True)
    risk.set_categories(Reference(ws, min_col=16, min_row=37, max_row=41))
    ws.add_chart(risk, "H35")

    product = BarChart()
    product.title = "品项销售额"
    product.height = 7
    product.width = 14
    product.add_data(Reference(ws, min_col=17, min_row=45, max_col=18, max_row=50), titles_from_data=True)
    product.set_categories(Reference(ws, min_col=16, min_row=46, max_row=50))
    ws.add_chart(product, "B51")


if __name__ == "__main__":
    build()
