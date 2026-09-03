# -*- coding: utf-8 -*-
from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


SOURCE_PATH = Path(r"E:\cxy\07-数据与指标\outputs\KA销售新增统计字段模板-大盘出货周转.xlsx")
OUTPUT_PATH = Path(r"E:\cxy\07-数据与指标\outputs\KA大盘出货数据-四页主题表.xlsx")

COLORS = {
    "ink": "17202A",
    "muted": "5B6675",
    "line": "D8DEE7",
    "bg": "F6F8FB",
    "panel": "FFFFFF",
    "header": "D9EAF7",
    "section": "EAF0F8",
    "blue": "1D4ED8",
    "green": "15803D",
    "amber": "B45309",
    "red": "B91C1C",
    "purple": "6D28D9",
}

thin = Side(style="thin", color=COLORS["line"])
border = Border(left=thin, right=thin, top=thin, bottom=thin)

SALES = "'销售明细'"
INV = "'库存明细'"
SHIP = "'出货退货'"


def prepare_sheet(ws, title: str, subtitle: str) -> None:
    ws.sheet_view.showGridLines = False
    for col in range(1, 22):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 16
    ws.column_dimensions["I"].width = 16
    ws.column_dimensions["J"].width = 16
    for row in range(1, 90):
        ws.row_dimensions[row].height = 22
    for row in ws["A1:U90"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=COLORS["bg"])
    ws.merge_cells("A1:U1")
    ws["A1"] = title
    ws["A1"].font = Font(name="微软雅黑", bold=True, size=18, color=COLORS["ink"])
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("A2:U2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="微软雅黑", size=9, color=COLORS["muted"])
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")


def style_range(ws, cell_range: str, fill: str = "FFFFFF") -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = border
            cell.font = Font(name="微软雅黑", size=10, color=COLORS["ink"])
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def section(ws, cell_range: str, title: str) -> None:
    ws.merge_cells(cell_range)
    style_range(ws, cell_range, COLORS["section"])
    cell = ws[cell_range.split(":")[0]]
    cell.value = title
    cell.font = Font(name="微软雅黑", bold=True, size=12, color=COLORS["ink"])
    cell.alignment = Alignment(horizontal="left", vertical="center")


def header(ws, row: int, start_col: int, titles: list[str]) -> None:
    for i, title in enumerate(titles, start_col):
        cell = ws.cell(row=row, column=i)
        cell.value = title
        cell.fill = PatternFill("solid", fgColor=COLORS["header"])
        cell.font = Font(name="微软雅黑", bold=True, size=10, color=COLORS["ink"])
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_validation(ws, cell: str, values: list[str]) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[cell])


def add_filters(ws, default_date_formula: str = "=MAX('销售明细'!A2:A500)") -> None:
    labels = [("B3", "统计日期", "C3", default_date_formula), ("E3", "KA系统", "F3", "全部"), ("H3", "大区", "I3", "全部"), ("K3", "省份", "L3", "全部")]
    for label_cell, label, value_cell, value in labels:
        ws[label_cell] = label
        ws[label_cell].font = Font(name="微软雅黑", bold=True, size=9, color=COLORS["muted"])
        ws[value_cell] = value
        ws[value_cell].fill = PatternFill("solid", fgColor=COLORS["panel"])
        ws[value_cell].border = border
        ws[value_cell].font = Font(name="微软雅黑", bold=True, size=10, color=COLORS["ink"])
        ws[value_cell].alignment = Alignment(horizontal="center", vertical="center")
    ws["C3"].number_format = "yyyy-mm-dd"
    add_validation(ws, "F3", ["全部", "大润发", "盒马", "华润", "沃尔玛", "永辉", "天虹", "比优特", "家得福"])
    add_validation(ws, "I3", ["全部", "华东", "华南", "华北", "华中", "西北", "东北"])
    add_validation(ws, "L3", ["全部", "上海", "江苏", "广东", "北京", "湖北", "浙江", "黑龙江", "天津"])


def opt(sheet: str, col: str, param: str) -> str:
    return f'--((({param}="全部")+({sheet}!${col}$2:${col}$500={param}))>0)'


def sales_cond(date_ref="$C$3") -> str:
    return f"({SALES}!$A$2:$A$500={date_ref})*{opt(SALES, 'B', '$F$3')}*{opt(SALES, 'C', '$I$3')}*{opt(SALES, 'D', '$L$3')}"


def ship_cond(date_ref="$C$3") -> str:
    return f"({SHIP}!$A$2:$A$500={date_ref})*{opt(SHIP, 'B', '$F$3')}*{opt(SHIP, 'C', '$I$3')}*{opt(SHIP, 'D', '$L$3')}"


def inv_cond(date_ref="$C$3") -> str:
    return f"({INV}!$A$2:$A$500={date_ref})*{opt(INV, 'B', '$F$3')}*{opt(INV, 'C', '$I$3')}*{opt(INV, 'D', '$L$3')}"


def sp(condition: str, sheet: str, col: str) -> str:
    return f"=SUMPRODUCT({condition}*{sheet}!${col}$2:${col}$500)"


def set_number(ws, cell_range: str, fmt: str) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.number_format = fmt


def copy_source_sheets(out_wb: Workbook) -> None:
    src = load_workbook(SOURCE_PATH, data_only=False)
    for source_name in ["销售明细", "库存明细", "出货退货"]:
        src_ws = src[source_name]
        ws = out_wb.create_sheet(source_name)
        for row in src_ws.iter_rows():
            for cell in row:
                new = ws[cell.coordinate]
                new.value = cell.value
                if cell.has_style:
                    new._style = copy(cell._style)
                new.number_format = cell.number_format
                new.alignment = copy(cell.alignment)
                new.font = copy(cell.font)
                new.fill = copy(cell.fill)
                new.border = copy(cell.border)
        for idx in range(1, src_ws.max_column + 1):
            ws.column_dimensions[get_column_letter(idx)].width = src_ws.column_dimensions[get_column_letter(idx)].width or 14
        ws.freeze_panes = "A2"
    src.close()


def build_shipment_sheet(ws) -> None:
    prepare_sheet(ws, "KA大盘出货数据", "关键词：系统每天出货数量、出货金额；by单品、by渠道、by数量、by金额。")
    add_filters(ws)
    section(ws, "B5:J5", "汇总：系统每天出货数量、出货金额（by系统/by数量/by金额）")
    header(ws, 6, 2, ["KA系统", "出货数量", "出货金额", "POS销售数量", "POS销售额", "出货/POS额", "退货数量", "退货金额", "数据状态"])
    systems = ["大润发", "盒马", "华润", "沃尔玛", "永辉", "天虹", "比优特", "家得福"]
    for r, system in enumerate(systems, 7):
        ws[f"B{r}"] = system
        sc = f"({SHIP}!$A$2:$A$500=$C$3)*--(({SHIP}!$B$2:$B$500=B{r})>0)*{opt(SHIP, 'C', '$I$3')}*{opt(SHIP, 'D', '$L$3')}"
        pc = f"({SALES}!$A$2:$A$500=$C$3)*--(({SALES}!$B$2:$B$500=B{r})>0)*{opt(SALES, 'C', '$I$3')}*{opt(SALES, 'D', '$L$3')}"
        ws[f"C{r}"] = sp(sc, SHIP, "I")
        ws[f"D{r}"] = sp(sc, SHIP, "J")
        ws[f"E{r}"] = sp(pc, SALES, "N")
        ws[f"F{r}"] = sp(pc, SALES, "O")
        ws[f"G{r}"] = f'=IFERROR(D{r}/F{r},0)'
        ws[f"H{r}"] = sp(sc, SHIP, "K")
        ws[f"I{r}"] = sp(sc, SHIP, "L")
        ws[f"J{r}"] = f'=IF(D{r}=0,"缺出货金额",IF(C{r}=0,"缺出货数量","可用"))'
    style_range(ws, "B7:J14")
    set_number(ws, "C7:F14", "#,##0.0")
    set_number(ws, "G7:G14", "0.0%")
    set_number(ws, "H7:I14", "#,##0.0")

    section(ws, "B17:J17", "汇总：by单品")
    header(ws, 18, 2, ["商品编码", "商品名称", "KA系统", "出货数量", "出货金额", "POS销售数量", "POS销售额", "出货/POS额", "备注"])
    skus = [
        ("4354168", "脱骨侠无骨鸡爪(柠檬味)", "大润发"),
        ("4502138", "脱骨侠虎皮凤爪(白松露味)", "大润发"),
        ("25354863", "常温脱骨侠无骨鸡爪柠檬", "沃尔玛"),
        ("431448153", "盒马工坊 蒜香无骨鸡爪 230g", "盒马"),
        ("v000110869", "脱骨侠双椒无骨鸭掌", "天虹"),
        ("10081319", "脱骨侠酸辣柠檬无骨鸡爪780g", "家得福"),
    ]
    for r, (sku, name, system) in enumerate(skus, 19):
        ws[f"B{r}"] = sku
        ws[f"C{r}"] = name
        ws[f"D{r}"] = system
        sc = f"{ship_cond()}*--(({SHIP}!$G$2:$G$500=B{r})>0)"
        pc = f"{sales_cond()}*--(({SALES}!$H$2:$H$500=B{r})>0)"
        ws[f"E{r}"] = sp(sc, SHIP, "I")
        ws[f"F{r}"] = sp(sc, SHIP, "J")
        ws[f"G{r}"] = sp(pc, SALES, "N")
        ws[f"H{r}"] = sp(pc, SALES, "O")
        ws[f"I{r}"] = f'=IFERROR(F{r}/H{r},0)'
        ws[f"J{r}"] = f'=IF(F{r}=0,"缺出货源","")'
    style_range(ws, "B19:J24")
    set_number(ws, "E19:H24", "#,##0.0")
    set_number(ws, "I19:I24", "0.0%")

    section(ws, "B27:J27", "汇总：by渠道")
    header(ws, 28, 2, ["渠道", "POS销售数量", "POS销售额", "出货数量", "出货金额", "出货/POS额", "说明"])
    channels = ["POS", "线下", "京东到家", "美团", "天虹到家"]
    for r, channel in enumerate(channels, 29):
        ws[f"B{r}"] = channel
        pc = f"{sales_cond()}*--(({SALES}!$M$2:$M$500=B{r})>0)"
        ws[f"C{r}"] = sp(pc, SALES, "N")
        ws[f"D{r}"] = sp(pc, SALES, "O")
        ws[f"E{r}"] = sp(ship_cond(), SHIP, "I")
        ws[f"F{r}"] = sp(ship_cond(), SHIP, "J")
        ws[f"G{r}"] = f'=IFERROR(F{r}/D{r},0)'
        ws[f"H{r}"] = "渠道出货需出货源提供渠道字段；当前出货按筛选总额展示"
    style_range(ws, "B29:H33")
    set_number(ws, "C29:F33", "#,##0.0")
    set_number(ws, "G29:G33", "0.0%")

    section(ws, "B36:J36", "原始数据：出货退货明细")
    header(ws, 37, 2, ["数据日期", "KA系统", "大区", "省份", "门店编码", "门店名称", "商品编码", "商品名称", "出货数量", "出货金额", "退货数量", "退货金额", "来源"])
    for i in range(2, 22):
        r = 36 + i
        ws[f"B{r}"] = f"='出货退货'!A{i}"
        ws[f"C{r}"] = f"='出货退货'!B{i}"
        ws[f"D{r}"] = f"='出货退货'!C{i}"
        ws[f"E{r}"] = f"='出货退货'!D{i}"
        ws[f"F{r}"] = f"='出货退货'!E{i}"
        ws[f"G{r}"] = f"='出货退货'!F{i}"
        ws[f"H{r}"] = f"='出货退货'!G{i}"
        ws[f"I{r}"] = f"='出货退货'!H{i}"
        ws[f"J{r}"] = f"='出货退货'!I{i}"
        ws[f"K{r}"] = f"='出货退货'!J{i}"
        ws[f"L{r}"] = f"='出货退货'!K{i}"
        ws[f"M{r}"] = f"='出货退货'!L{i}"
        ws[f"N{r}"] = f"='出货退货'!M{i}"
    style_range(ws, "B38:N56")
    set_number(ws, "J38:M56", "#,##0.0")
    add_shipment_charts(ws)


def add_shipment_charts(ws) -> None:
    chart = BarChart()
    chart.title = "系统出货金额 vs POS销售额"
    chart.y_axis.title = "金额"
    chart.height = 8
    chart.width = 15
    chart.add_data(Reference(ws, min_col=4, min_row=6, max_col=6, max_row=14), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=7, max_row=14))
    ws.add_chart(chart, "L5")
    chart2 = BarChart()
    chart2.title = "by单品出货数量"
    chart2.y_axis.title = "数量"
    chart2.height = 8
    chart2.width = 15
    chart2.add_data(Reference(ws, min_col=5, min_row=18, max_row=24), titles_from_data=True)
    chart2.set_categories(Reference(ws, min_col=3, min_row=19, max_row=24))
    ws.add_chart(chart2, "L21")


def build_compare_sheet(ws) -> None:
    prepare_sheet(ws, "系统单品出货额和POS额对比", "关键词：系统、单品、出货额、POS额、差异、比例。")
    add_filters(ws)
    section(ws, "B5:J5", "汇总：系统单品出货额和POS额对比")
    header(ws, 6, 2, ["KA系统", "商品编码", "商品名称", "出货额", "POS额", "差异", "出货/POS额", "出货数量", "POS销量", "状态"])
    rows = [
        ("大润发", "4354168", "脱骨侠无骨鸡爪(柠檬味)"),
        ("大润发", "4502138", "脱骨侠虎皮凤爪(白松露味)"),
        ("沃尔玛", "25354863", "常温脱骨侠无骨鸡爪柠檬"),
        ("盒马", "431448153", "盒马工坊 蒜香无骨鸡爪 230g"),
        ("天虹", "v000110869", "脱骨侠双椒无骨鸭掌"),
        ("家得福", "10081319", "脱骨侠酸辣柠檬无骨鸡爪780g"),
    ]
    for r, (system, sku, name) in enumerate(rows, 7):
        ws[f"B{r}"] = system
        ws[f"C{r}"] = sku
        ws[f"D{r}"] = name
        sc = f"{ship_cond()}*--(({SHIP}!$B$2:$B$500=B{r})>0)*--(({SHIP}!$G$2:$G$500=C{r})>0)"
        pc = f"{sales_cond()}*--(({SALES}!$B$2:$B$500=B{r})>0)*--(({SALES}!$H$2:$H$500=C{r})>0)"
        ws[f"E{r}"] = sp(sc, SHIP, "J")
        ws[f"F{r}"] = sp(pc, SALES, "O")
        ws[f"G{r}"] = f"=E{r}-F{r}"
        ws[f"H{r}"] = f'=IFERROR(E{r}/F{r},0)'
        ws[f"I{r}"] = sp(sc, SHIP, "I")
        ws[f"J{r}"] = sp(pc, SALES, "N")
        ws[f"K{r}"] = f'=IF(E{r}=0,"缺出货源",IF(H{r}<0.8,"出货低于POS",IF(H{r}>1.2,"出货高于POS","匹配")))'
    style_range(ws, "B7:K12")
    set_number(ws, "E7:G12", "#,##0.0")
    set_number(ws, "H7:H12", "0.0%")
    set_number(ws, "I7:J12", "#,##0.0")
    section(ws, "B15:K15", "原始数据：POS销售明细与出货明细")
    header(ws, 16, 2, ["来源", "数据日期", "KA系统", "商品编码", "商品名称", "门店名称", "数量", "金额", "渠道/来源字段"])
    for i in range(2, 17):
        r = 16 + i
        ws[f"B{r}"] = "POS"
        ws[f"C{r}"] = f"='销售明细'!A{i}"
        ws[f"D{r}"] = f"='销售明细'!B{i}"
        ws[f"E{r}"] = f"='销售明细'!H{i}"
        ws[f"F{r}"] = f"='销售明细'!J{i}"
        ws[f"G{r}"] = f"='销售明细'!F{i}"
        ws[f"H{r}"] = f"='销售明细'!N{i}"
        ws[f"I{r}"] = f"='销售明细'!O{i}"
        ws[f"J{r}"] = f"='销售明细'!M{i}"
    for i in range(2, 7):
        r = 33 + i
        ws[f"B{r}"] = "出货"
        ws[f"C{r}"] = f"='出货退货'!A{i}"
        ws[f"D{r}"] = f"='出货退货'!B{i}"
        ws[f"E{r}"] = f"='出货退货'!G{i}"
        ws[f"F{r}"] = f"='出货退货'!H{i}"
        ws[f"G{r}"] = f"='出货退货'!F{i}"
        ws[f"H{r}"] = f"='出货退货'!I{i}"
        ws[f"I{r}"] = f"='出货退货'!J{i}"
        ws[f"J{r}"] = f"='出货退货'!M{i}"
    style_range(ws, "B18:J40")
    set_number(ws, "H18:I40", "#,##0.0")
    chart = BarChart()
    chart.title = "单品出货额 vs POS额"
    chart.height = 8
    chart.width = 16
    chart.add_data(Reference(ws, min_col=5, min_row=6, max_col=6, max_row=12), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=4, min_row=7, max_row=12))
    ws.add_chart(chart, "M5")
    chart2 = BarChart()
    chart2.title = "出货/POS额比例"
    chart2.height = 8
    chart2.width = 16
    chart2.add_data(Reference(ws, min_col=8, min_row=6, max_row=12), titles_from_data=True)
    chart2.set_categories(Reference(ws, min_col=4, min_row=7, max_row=12))
    ws.add_chart(chart2, "M22")


def build_turnover_sheet(ws) -> None:
    prepare_sheet(ws, "门店端库存周转天数", "关键词：POS数据、门店端库存、各大区、各单品、数据日前7天平均。")
    add_filters(ws)
    section(ws, "B5:J5", "汇总：各大区区分各单品库存周转天数")
    header(ws, 6, 2, ["大区", "商品编码", "商品名称", "门店库存", "近7日平均销量", "库存周转天数", "有库存门店数", "动销门店数", "风险"])
    rows = [
        ("华东", "4354168", "脱骨侠无骨鸡爪(柠檬味)"),
        ("华南", "4502138", "脱骨侠虎皮凤爪(白松露味)"),
        ("华南", "25354863", "常温脱骨侠无骨鸡爪柠檬"),
        ("华北", "5150673", "脱骨侠无骨鸡爪(柠檬味)158g"),
        ("华中", "431448153", "盒马工坊 蒜香无骨鸡爪 230g"),
        ("华东", "10081319", "脱骨侠酸辣柠檬无骨鸡爪780g"),
    ]
    for r, (region, sku, name) in enumerate(rows, 7):
        ws[f"B{r}"] = region
        ws[f"C{r}"] = sku
        ws[f"D{r}"] = name
        cond = f"{inv_cond()}*--(({INV}!$C$2:$C$500=B{r})>0)*--(({INV}!$H$2:$H$500=C{r})>0)"
        ws[f"E{r}"] = sp(cond, INV, "J")
        ws[f"F{r}"] = sp(cond, INV, "M")
        ws[f"G{r}"] = f'=IFERROR(E{r}/F{r},"")'
        ws[f"H{r}"] = f'=COUNTIFS({INV}!$A$2:$A$500,$C$3,{INV}!$C$2:$C$500,B{r},{INV}!$H$2:$H$500,C{r},{INV}!$J$2:$J$500,">0")'
        ws[f"I{r}"] = f'=COUNTIFS({SALES}!$A$2:$A$500,$C$3,{SALES}!$C$2:$C$500,B{r},{SALES}!$H$2:$H$500,C{r},{SALES}!$N$2:$N$500,">0")'
        ws[f"J{r}"] = f'=IF(G{r}="","无动销",IF(G{r}>60,"高风险","正常"))'
    style_range(ws, "B7:J12")
    set_number(ws, "E7:I12", "#,##0.0")
    ws.conditional_formatting.add("G7:G12", CellIsRule(operator="greaterThan", formula=["60"], fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(color=COLORS["red"], bold=True)))
    ws.conditional_formatting.add("G7:G12", IconSetRule("3TrafficLights1", "num", [30, 60, 120], showValue=True))
    section(ws, "B15:J15", "原始数据：库存明细")
    header(ws, 16, 2, ["数据日期", "KA系统", "大区", "省份", "门店编码", "门店名称", "商品编码", "商品名称", "门店库存", "近7日平均销量", "周转天数", "是否超60天"])
    for i in range(2, 31):
        r = 16 + i
        for col, src_col in enumerate(range(1, 16), 2):
            # Skip DC库存/在途列 for this focused raw view.
            pass
        mappings = [("B", "A"), ("C", "B"), ("D", "C"), ("E", "D"), ("F", "E"), ("G", "F"), ("H", "H"), ("I", "I"), ("J", "J"), ("K", "M"), ("L", "N"), ("M", "O")]
        for dest, src_col in mappings:
            ws[f"{dest}{r}"] = f"='库存明细'!{src_col}{i}"
    style_range(ws, "B18:M45")
    set_number(ws, "J18:L45", "#,##0.0")
    chart = BarChart()
    chart.title = "各大区单品库存周转天数"
    chart.height = 8
    chart.width = 16
    chart.add_data(Reference(ws, min_col=7, min_row=6, max_row=12), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=4, min_row=7, max_row=12))
    ws.add_chart(chart, "M5")
    chart2 = LineChart()
    chart2.title = "门店库存 vs 近7日均销"
    chart2.height = 8
    chart2.width = 16
    chart2.add_data(Reference(ws, min_col=5, min_row=6, max_col=6, max_row=12), titles_from_data=True)
    chart2.set_categories(Reference(ws, min_col=4, min_row=7, max_row=12))
    ws.add_chart(chart2, "M22")


def build_over60_sheet(ws) -> None:
    prepare_sheet(ws, "超60天门店排名", "关键词：各大区、库存周转天数超60天、门店数排名。")
    add_filters(ws)
    section(ws, "B5:J5", "汇总：各大区库存周转天数超60天门店数排名")
    header(ws, 6, 2, ["排名", "大区", "超60天门店数", "门店库存", "近7日平均销量", "平均周转天数", "主要风险单品", "处理建议"])
    regions = ["华东", "华南", "华北", "华中", "东北"]
    for idx, region in enumerate(regions, 20):
        ws[f"L{idx}"] = region
        cond = f"{inv_cond()}*--(({INV}!$C$2:$C$500=L{idx})>0)"
        ws[f"M{idx}"] = f'=COUNTIFS({INV}!$A$2:$A$500,$C$3,{INV}!$C$2:$C$500,L{idx},{INV}!$O$2:$O$500,"Y")'
        ws[f"N{idx}"] = sp(cond, INV, "J")
        ws[f"O{idx}"] = sp(cond, INV, "M")
        ws[f"P{idx}"] = f'=IFERROR(N{idx}/O{idx},"")'
    for i in range(5):
        r = 7 + i
        ws[f"B{r}"] = i + 1
        ws[f"C{r}"] = f'=IFERROR(INDEX($L$20:$L$24,MATCH(LARGE($M$20:$M$24,B{r}),$M$20:$M$24,0)),"")'
        ws[f"D{r}"] = f'=IFERROR(INDEX($M$20:$M$24,MATCH(C{r},$L$20:$L$24,0)),0)'
        ws[f"E{r}"] = f'=IFERROR(INDEX($N$20:$N$24,MATCH(C{r},$L$20:$L$24,0)),0)'
        ws[f"F{r}"] = f'=IFERROR(INDEX($O$20:$O$24,MATCH(C{r},$L$20:$L$24,0)),0)'
        ws[f"G{r}"] = f'=IFERROR(INDEX($P$20:$P$24,MATCH(C{r},$L$20:$L$24,0)),"")'
        ws[f"H{r}"] = '=IFERROR(INDEX(\'库存明细\'!$I$2:$I$500,MATCH(MAXIFS(\'库存明细\'!$N$2:$N$500,\'库存明细\'!$C$2:$C$500,C{0}),\'库存明细\'!$N$2:$N$500,0)),"")'.format(r)
        ws[f"I{r}"] = f'=IF(D{r}>0,"优先清理无动销库存/调拨/促销","正常跟踪")'
    style_range(ws, "B7:I11")
    set_number(ws, "D7:G11", "#,##0.0")
    ws.conditional_formatting.add("D7:D11", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(color=COLORS["red"], bold=True)))
    section(ws, "B15:J15", "原始数据：超60天门店明细")
    header(ws, 16, 2, ["数据日期", "KA系统", "大区", "省份", "门店编码", "门店名称", "商品编码", "商品名称", "门店库存", "近7日平均销量", "周转天数", "是否超60天"])
    for i in range(2, 31):
        r = 16 + i
        mappings = [("B", "A"), ("C", "B"), ("D", "C"), ("E", "D"), ("F", "E"), ("G", "F"), ("H", "H"), ("I", "I"), ("J", "J"), ("K", "M"), ("L", "N"), ("M", "O")]
        for dest, src_col in mappings:
            ws[f"{dest}{r}"] = f"='库存明细'!{src_col}{i}"
    style_range(ws, "B18:M45")
    set_number(ws, "J18:L45", "#,##0.0")
    chart = BarChart()
    chart.type = "bar"
    chart.title = "各大区超60天门店数排名"
    chart.height = 8
    chart.width = 16
    chart.add_data(Reference(ws, min_col=4, min_row=6, max_row=11), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=3, min_row=7, max_row=11))
    ws.add_chart(chart, "M5")
    chart2 = BarChart()
    chart2.title = "区域平均周转天数"
    chart2.height = 8
    chart2.width = 16
    chart2.add_data(Reference(ws, min_col=7, min_row=6, max_row=11), titles_from_data=True)
    chart2.set_categories(Reference(ws, min_col=3, min_row=7, max_row=11))
    ws.add_chart(chart2, "M22")


def build() -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    build_shipment_sheet(wb.create_sheet("1-KA大盘出货数据"))
    build_compare_sheet(wb.create_sheet("2-单品出货POS对比"))
    build_turnover_sheet(wb.create_sheet("3-门店库存周转天数"))
    build_over60_sheet(wb.create_sheet("4-超60天门店排名"))
    copy_source_sheets(wb)
    for ws in wb.worksheets:
        ws.freeze_panes = "A6" if ws.title.startswith(("1-", "2-", "3-", "4-")) else "A2"
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    check = load_workbook(OUTPUT_PATH, data_only=False)
    expected = ["1-KA大盘出货数据", "2-单品出货POS对比", "3-门店库存周转天数", "4-超60天门店排名", "销售明细", "库存明细", "出货退货"]
    assert check.sheetnames == expected
    assert len(check["1-KA大盘出货数据"]._charts) == 2
    assert len(check["2-单品出货POS对比"]._charts) == 2
    assert len(check["3-门店库存周转天数"]._charts) == 2
    assert len(check["4-超60天门店排名"]._charts) == 2
    assert check["1-KA大盘出货数据"]["B5"].value.startswith("汇总")
    check.close()
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build()
