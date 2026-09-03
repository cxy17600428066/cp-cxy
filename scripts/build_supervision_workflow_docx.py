# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"E:/cxy/outputs/docx/督办全流程设计方案-领导汇报版.docx")


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False, fill=None, color=None, align=None):
    if fill:
        shade(cell, fill)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align or (WD_ALIGN_PARAGRAPH.CENTER if len(str(text)) <= 10 else WD_ALIGN_PARAGRAPH.LEFT)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_font(run, size=9.5, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Light List Accent 1"
    for idx, header in enumerate(headers):
        set_cell(tbl.rows[0].cells[idx], header, bold=True, fill="0F766E", color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        row = tbl.add_row().cells
        for idx, value in enumerate(row_data):
            set_cell(row[idx], value)
    if widths:
        for row in tbl.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return tbl


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    r = p.add_run(text)
    set_font(r, size=16 if level == 1 else 12.5, bold=True, color=(15, 83, 76) if level == 1 else (31, 68, 103))
    return p


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, 10.5)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_font(r, 10.5)


def callout(doc, title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Light Shading Accent 1"
    cell = tbl.cell(0, 0)
    shade(cell, "E8F5F3")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, size=11, bold=True, color=(15, 83, 76))
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(text)
    set_font(r2, size=10.5)
    doc.add_paragraph()


def cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(110)
    r = p.add_run("督办全流程设计方案")
    set_font(r, size=28, bold=True, color=(15, 83, 76))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("领导汇报版")
    set_font(r, size=16, color=(90, 105, 120))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("覆盖个人发起、领导发起、部门发起三类督办场景")
    set_font(r, size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(150)
    r = p.add_run("2026年6月")
    set_font(r, size=11, color=(90, 105, 120))
    doc.add_page_break()


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[name].font.name = "Microsoft YaHei"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    cover(doc)

    heading(doc, "一、方案摘要", 1)
    callout(
        doc,
        "建设目标",
        "建立一套覆盖发起、审批、派发、签收、执行、反馈、验收、评分、归档的督办闭环机制。重点解决事项派发后无人签收、责任人拒收、延期无审批、转办无闭环、催办无留痕、领导批示无法落地等管理痛点。",
    )
    body(
        doc,
        "本方案将督办事项作为独立的管理对象，不依赖某一个具体 OA 审批流程。无论事项来自个人、领导还是部门，只要进入督办体系，就必须具备责任人、完成时限、过程反馈、异常处理、亮灯预警、验收评分和归档记录。",
    )
    bullets(
        doc,
        [
            "个人发起：先审批确认督办是否成立，再派发责任人，防止普通协作事项泛化为正式督办。",
            "领导发起：可以跳过普通审批，但必须经过签收、执行反馈、验收评分和归档。",
            "部门发起：支持本部门派发和跨部门承接，跨部门拒绝必须进入协调或领导批示。",
            "所有异常动作都必须回到流程主线，不能形成悬空状态。",
        ],
    )

    heading(doc, "二、三类督办流程概览", 1)
    table(
        doc,
        ["流程类型", "适用场景", "关键控制点", "闭环要求"],
        [
            ["个人发起督办", "员工发现跨部门事项、长期未推动事项、本人事项需正式跟踪", "发起后必须审批；审批通过后才允许派发", "验收通过并评分后办结"],
            ["领导发起督办", "领导指示、会议纪要、重点专项、临时交办事项", "可直接生成督办；责任人必须签收；拒收需抄送领导", "领导或督办员验收评分后归档"],
            ["部门发起督办", "部门间协作、制度执行、项目推进、专项整改", "部门负责人确认；跨部门需接收确认", "跨部门拒绝进入协调或领导批示"],
        ],
        widths=[1.2, 2.05, 2.15, 1.75],
    )

    heading(doc, "三、总体流转主线", 1)
    body(doc, "督办主线统一为 9 个阶段。不同发起类型可以在前置环节有差异，但进入派发后应使用统一的执行、验收和归档规则。")
    table(
        doc,
        ["阶段", "核心动作", "责任角色", "关键规则"],
        [
            ["1. 发起", "填写督办事项、目标、时限、附件", "发起人 / 领导 / 部门", "信息不完整不能提交"],
            ["2. 审批或确认", "判断督办是否成立", "审批人 / 部门负责人", "不成立必须说明原因"],
            ["3. 派发", "指定责任人、协办人、完成时限", "派发人 / 督办员", "没有责任人和时限不能派发"],
            ["4. 签收", "责任人确认接收任务", "责任人", "不签收必须催签并升级"],
            ["5. 执行反馈", "提交进度、卡点、附件、完成结果", "责任人 / 协办人", "长期无反馈自动催办"],
            ["6. 异常处理", "延期、变更、转办、拒收、批示", "相关责任角色", "异常不能悬空，必须回到主线"],
            ["7. 验收", "确认完成质量和交付物", "发起人 / 督办人 / 领导", "退回必须写整改要求"],
            ["8. 评分", "按时效、质量、协同、材料完整性评分", "验收人 / 领导", "评分后才允许办结"],
            ["9. 归档", "形成结论和过程留痕", "督办人 / 系统", "归档后进入统计分析"],
        ],
        widths=[1.05, 2.05, 1.55, 2.5],
    )

    heading(doc, "四、关键闭环机制", 1)
    heading(doc, "1. 签收闭环", 2)
    body(doc, "派发后必须进入签收环节。责任人签收时需要确认任务范围、完成标准、完成时限和交付物。未签收、延期签收、拒收都必须进入明确处理路径。")
    table(
        doc,
        ["情形", "系统动作", "处理结果"],
        [
            ["正常签收", "记录签收时间和签收人", "进入执行中"],
            ["4 小时未签收", "提醒责任人", "保持待签收"],
            ["8 小时未签收", "抄送派发人并亮黄灯", "进入签收预警"],
            ["1 个工作日未签收", "抄送责任人上级并亮红灯", "进入签收超期"],
            ["仍未签收", "要求派发人改派、强制签收或升级批示", "进入升级处理中"],
            ["责任人拒收", "填写拒收原因和建议责任人", "派发人同意后重派，驳回后继续签收"],
        ],
        widths=[1.35, 3.05, 2.65],
    )

    heading(doc, "2. 异常闭环", 2)
    table(
        doc,
        ["异常动作", "必须填写内容", "闭环规则"],
        [
            ["延期", "延期原因、新完成时间、影响范围", "同意后更新时限并记录延期次数；驳回则维持原时限"],
            ["变更", "变更内容、变更原因、影响说明", "生成变更版本，不能覆盖原督办要求"],
            ["转办", "转办原因、新责任人、责任说明", "新责任人必须重新签收"],
            ["拒收", "拒收原因、建议责任人、是否需协调", "不能直接关闭，必须回到派发或领导批示"],
            ["批示", "批示意见、承接动作、责任人", "批示必须形成新待办或推动状态变化"],
            ["催办", "催办对象、催办内容、催办时间", "必须留痕并进入统计"],
        ],
        widths=[1.1, 2.75, 3.2],
    )

    heading(doc, "五、亮灯预警规则", 1)
    body(doc, "亮灯规则用于让领导和督办人员快速识别风险。建议将签收、进度、到期、验收分别亮灯，总体灯色取最严重状态。")
    table(
        doc,
        ["灯色", "触发条件", "系统动作"],
        [
            ["绿灯", "未临期，正常反馈", "正常展示"],
            ["黄灯", "距截止时间小于 20%，或签收/反馈即将超时", "提醒责任人"],
            ["红灯", "已超期，或拒收/延期/转办未处理", "抄送上级和督办人"],
            ["黑灯", "严重超期，或重大事项无人处理", "升级领导批示"],
        ],
        widths=[1.0, 3.15, 2.9],
    )

    heading(doc, "六、验收评分机制", 1)
    body(doc, "督办事项完成后不应直接关闭，必须先验收，再评分，最后归档。评分结果可用于部门执行力分析和责任人绩效参考。")
    table(
        doc,
        ["评分项", "分值", "说明"],
        [
            ["按时完成", "30", "是否按计划完成，是否发生超期"],
            ["完成质量", "30", "是否达到督办目标和交付标准"],
            ["反馈及时性", "20", "是否按要求反馈进展和卡点"],
            ["协同配合", "10", "是否主动协同相关部门"],
            ["材料完整性", "10", "附件、说明、过程记录是否完整"],
        ],
        widths=[1.4, 0.75, 4.9],
    )
    body(doc, "评分建议：90-100 为优秀，80-89 为良好，60-79 为合格，60 分以下为不合格并进入复盘。")

    heading(doc, "七、领导关注看板", 1)
    body(doc, "领导端不需要处理所有细节，应聚焦重大事项、超期事项、拒收事项和跨部门争议。建议建设以下看板指标。")
    bullets(
        doc,
        [
            "督办总量、待签收数量、执行中数量、已超期数量。",
            "按部门、责任人、灯色统计的风险分布。",
            "重大督办事项清单和黑灯事项清单。",
            "拒收、延期、转办、变更次数排行。",
            "平均办结周期、按时完成率、验收通过率、评分分布。",
        ],
    )

    heading(doc, "八、落地建议", 1)
    table(
        doc,
        ["阶段", "建设重点", "交付结果"],
        [
            ["第一阶段", "统一督办单、状态、签收、催办、亮灯规则", "形成可运行的基础督办闭环"],
            ["第二阶段", "补齐延期、变更、转办、拒收、批示机制", "解决异常事项悬空问题"],
            ["第三阶段", "上线验收评分和领导看板", "形成执行力分析和管理抓手"],
            ["第四阶段", "接入会议纪要、项目计划、重点工作来源", "扩展为统一事项督办平台"],
        ],
        widths=[1.1, 3.3, 2.65],
    )

    callout(doc, "结论", "督办系统的关键不是“能发起”，而是“必须闭环”。建议将签收、拒收、延期、变更、转办、催办、批示、验收、评分全部设计为可追踪、可升级、可统计的流程动作，确保事项不会停在某个无人处理的状态。")

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("督办全流程设计方案｜领导汇报版")
        set_font(run, size=9, color=(120, 130, 140))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
