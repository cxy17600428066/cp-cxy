# -*- coding: utf-8 -*-
from __future__ import annotations

import math
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


OUTPUT = Path(r"E:/cxy/outputs/docx/督办全流程设计方案-领导汇报版-含流程图.docx")
VERIFY_COPY = Path(r"E:/cxy/outputs/docx/supervision_workflow_leadership_with_charts.docx")
CHART_DIR = Path(r"E:/cxy/outputs/docx/flowcharts")

TEAL = (15, 118, 110)
TEAL_DARK = (15, 83, 76)
TEAL_LIGHT = (232, 245, 243)
BLUE = (37, 99, 235)
BLUE_LIGHT = (230, 240, 255)
AMBER = (217, 119, 6)
AMBER_LIGHT = (255, 247, 237)
RED = (190, 18, 60)
RED_LIGHT = (255, 241, 242)
GRAY = (71, 85, 105)
GRAY_LIGHT = (248, 250, 252)
LINE = (51, 65, 85)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        r"C:/Windows/Fonts/NotoSansSC-VF.ttf",
        r"C:/Windows/Fonts/simhei.ttf" if bold else r"C:/Windows/Fonts/Deng.ttf",
        r"C:/Windows/Fonts/simsun.ttc",
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False, fill=None, color=None, align=None):
    if fill:
        shade(cell, fill)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align or (
        WD_ALIGN_PARAGRAPH.CENTER if len(str(text)) <= 10 else WD_ALIGN_PARAGRAPH.LEFT
    )
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(str(text))
    set_font(run, size=9.2, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Light List Accent 1"
    for idx, header in enumerate(headers):
        set_cell(
            tbl.rows[0].cells[idx],
            header,
            bold=True,
            fill="0F766E",
            color=(255, 255, 255),
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row_data in rows:
        row = tbl.add_row().cells
        for idx, value in enumerate(row_data):
            set_cell(row[idx], value)
    if widths:
        for row in tbl.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return tbl


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    r = p.add_run(text)
    set_font(
        r,
        size=16 if level == 1 else 12.5,
        bold=True,
        color=(15, 83, 76) if level == 1 else (31, 68, 103),
    )
    return p


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, 10.5)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_font(r, 10.5)


def callout(doc, title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Light Shading Accent 1"
    cell = tbl.cell(0, 0)
    shade(cell, "E8F5F3")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, size=11, bold=True, color=(15, 83, 76))
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(text)
    set_font(r2, size=10.5)
    doc.add_paragraph()


def cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(98)
    r = p.add_run("督办全流程设计方案")
    set_font(r, size=28, bold=True, color=(15, 83, 76))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("领导汇报版｜含流程图")
    set_font(r, size=16, color=(90, 105, 120))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("覆盖个人发起、领导发起、部门发起三类督办场景")
    set_font(r, size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(150)
    r = p.add_run("2026年6月")
    set_font(r, size=11, color=(90, 105, 120))
    doc.add_page_break()


def wrap_text(draw, text, face, max_width):
    lines = []
    for part in str(text).split("\n"):
        line = ""
        for char in part:
            candidate = line + char
            if draw.textlength(candidate, font=face) <= max_width or not line:
                line = candidate
            else:
                lines.append(line)
                line = char
        if line:
            lines.append(line)
    return lines or [""]


def draw_text_center(draw, box, text, face, fill=(15, 23, 42), line_gap=5):
    x, y, w, h = box
    lines = wrap_text(draw, text, face, w - 28)
    heights = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=face)
        heights.append(bbox[3] - bbox[1])
    total_height = sum(heights) + line_gap * (len(lines) - 1)
    cur_y = y + (h - total_height) / 2
    for line, line_h in zip(lines, heights):
        line_w = draw.textlength(line, font=face)
        draw.text((x + (w - line_w) / 2, cur_y), line, font=face, fill=fill)
        cur_y += line_h + line_gap


def draw_node(draw, node):
    x, y, w, h = node["x"], node["y"], node["w"], node["h"]
    text = node["text"]
    kind = node.get("kind", "process")
    face = font(node.get("font_size", 27), node.get("bold", True))
    fill = node.get("fill", TEAL_LIGHT)
    outline = node.get("outline", TEAL)
    text_fill = node.get("text_fill", (15, 23, 42))
    if kind == "decision":
        points = [(x + w / 2, y), (x + w, y + h / 2), (x + w / 2, y + h), (x, y + h / 2)]
        draw.polygon(points, fill=fill, outline=outline)
        draw.line(points + [points[0]], fill=outline, width=4)
        draw_text_center(draw, (x + 12, y + 8, w - 24, h - 16), text, face, text_fill)
    else:
        radius = 26 if kind != "small" else 18
        draw.rounded_rectangle((x, y, x + w, y + h), radius=radius, fill=fill, outline=outline, width=4)
        draw_text_center(draw, (x, y, w, h), text, face, text_fill)


def anchor(node, side):
    x, y, w, h = node["x"], node["y"], node["w"], node["h"]
    if side == "left":
        return (x, y + h / 2)
    if side == "right":
        return (x + w, y + h / 2)
    if side == "top":
        return (x + w / 2, y)
    if side == "bottom":
        return (x + w / 2, y + h)
    return (x + w / 2, y + h / 2)


def draw_arrowhead(draw, p1, p2, color):
    angle = math.atan2(p2[1] - p1[1], p2[0] - p1[0])
    size = 16
    points = [
        p2,
        (p2[0] - size * math.cos(angle - math.pi / 6), p2[1] - size * math.sin(angle - math.pi / 6)),
        (p2[0] - size * math.cos(angle + math.pi / 6), p2[1] - size * math.sin(angle + math.pi / 6)),
    ]
    draw.polygon(points, fill=color)


def label(draw, x, y, text):
    if not text:
        return
    face = font(22, True)
    padding_x, padding_y = 9, 5
    w = draw.textlength(text, font=face)
    bbox = draw.textbbox((0, 0), text, font=face)
    h = bbox[3] - bbox[1]
    draw.rounded_rectangle(
        (x - padding_x, y - padding_y, x + w + padding_x, y + h + padding_y),
        radius=10,
        fill=(255, 255, 255),
        outline=(226, 232, 240),
        width=2,
    )
    draw.text((x, y), text, font=face, fill=GRAY)


def draw_edge(draw, nodes, start, end, start_side="right", end_side="left", text="", via=None, color=LINE):
    points = [anchor(nodes[start], start_side)] + (via or []) + [anchor(nodes[end], end_side)]
    draw.line(points, fill=color, width=5, joint="curve")
    draw_arrowhead(draw, points[-2], points[-1], color)


def new_canvas(title, subtitle="", size=(1800, 1000)):
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, size[0], 96), fill=TEAL_DARK)
    draw.text((54, 22), title, font=font(42, True), fill="white")
    if subtitle:
        draw.text((54, 106), subtitle, font=font(24), fill=GRAY)
    return image, draw


def save_diagram(image, filename):
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    path = CHART_DIR / filename
    image.save(path, dpi=(180, 180))
    return path


def diagram_overall():
    image, draw = new_canvas("督办总体流转主线", "所有督办事项最终都必须经过签收、执行、验收、评分和归档。", (1800, 960))
    labels = [
        "发起\n目标/时限/附件",
        "审批/确认\n是否成立",
        "派发\n责任人/协办人",
        "签收\n确认范围",
        "执行反馈\n进度/卡点",
        "异常处理\n延期/变更等",
        "验收\n确认交付",
        "评分\n时效/质量",
        "归档\n统计复盘",
    ]
    nodes = {}
    for i, item in enumerate(labels):
        nodes[f"n{i}"] = {
            "x": 50 + i * 194,
            "y": 215,
            "w": 155,
            "h": 112,
            "text": item,
            "font_size": 24,
        }
    for key in nodes:
        draw_node(draw, nodes[key])
    for i in range(len(labels) - 1):
        draw_edge(draw, nodes, f"n{i}", f"n{i+1}")

    smalls = ["延期", "变更", "转办", "拒收", "催办", "批示"]
    for i, item in enumerate(smalls):
        nodes[f"s{i}"] = {
            "x": 405 + i * 170,
            "y": 555,
            "w": 125,
            "h": 72,
            "text": item,
            "kind": "small",
            "font_size": 24,
            "fill": AMBER_LIGHT,
            "outline": AMBER,
        }
        draw_node(draw, nodes[f"s{i}"])
        draw_edge(draw, nodes, "n5", f"s{i}", "bottom", "top", color=AMBER)
        draw_edge(draw, nodes, f"s{i}", "n5", "bottom", "bottom", color=AMBER, via=[(nodes[f"s{i}"]["x"] + 62, 720), (nodes["n5"]["x"] + 77, 720)])

    callouts = [
        ("无人签收", "催签、亮灯、升级"),
        ("责任争议", "重派或领导批示"),
        ("完成质量不足", "验收退回整改"),
    ]
    for i, (title, desc) in enumerate(callouts):
        x = 260 + i * 430
        draw.rounded_rectangle((x, 790, x + 350, 875), radius=20, fill=GRAY_LIGHT, outline=(203, 213, 225), width=3)
        draw.text((x + 24, 805), title, font=font(25, True), fill=TEAL_DARK)
        draw.text((x + 24, 842), desc, font=font(22), fill=GRAY)
    return save_diagram(image, "01-overall-flow.png")


def diagram_personal_simple():
    image, draw = new_canvas("个人发起督办流程", "个人发起必须先审批，审批成立后才派发责任人；异常必须回到主线或关闭留痕。", (1800, 900))
    draw.text((70, 185), "主线", font=font(28, True), fill=TEAL_DARK)
    draw.text((70, 560), "异常闭环", font=font(28, True), fill=RED)
    nodes = {}
    main = [
        ("个人填写\n督办申请", "a"),
        ("审批确认\n是否成立", "b"),
        ("派发\n责任人/时限", "c"),
        ("责任人\n签收", "d"),
        ("执行反馈\n进度/结果", "e"),
        ("验收评分", "f"),
        ("办结归档", "g"),
    ]
    start_x, gap, w, h = 150, 42, 195, 102
    for idx, (text, key) in enumerate(main):
        nodes[key] = {
            "x": start_x + idx * (w + gap),
            "y": 250,
            "w": w,
            "h": h,
            "text": text,
            "font_size": 25,
            "fill": BLUE_LIGHT if key in {"d", "e", "f"} else TEAL_LIGHT,
            "outline": BLUE if key in {"d", "e", "f"} else TEAL,
        }
    abnormal = [
        ("审批退回\n补充材料", "r1", 240),
        ("不成立\n关闭并记录原因", "r2", 575),
        ("拒收/超时未签收\n重派或升级批示", "r3", 960),
        ("验收退回\n继续整改反馈", "r4", 1345),
    ]
    for text, key, x in abnormal:
        nodes[key] = {
            "x": x,
            "y": 620,
            "w": 260,
            "h": 98,
            "text": text,
            "font_size": 23,
            "fill": RED_LIGHT,
            "outline": RED,
        }
    for node in nodes.values():
        draw_node(draw, node)
    for left, right in zip(["a", "b", "c", "d", "e", "f"], ["b", "c", "d", "e", "f", "g"]):
        draw_edge(draw, nodes, left, right)
    draw_edge(draw, nodes, "b", "r1", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r1", "a", "top", "bottom", color=RED)
    draw_edge(draw, nodes, "b", "r2", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "d", "r3", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r3", "c", "top", "bottom", color=RED)
    draw_edge(draw, nodes, "f", "r4", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r4", "e", "top", "bottom", color=RED)
    draw.rounded_rectangle((108, 775, 1692, 835), radius=18, fill=GRAY_LIGHT, outline=(203, 213, 225), width=3)
    draw.text((140, 792), "控制点：个人发起先审批；派发后必须签收；拒收、超时、验收退回都不能停留，必须重派、升级或整改。", font=font(24, True), fill=GRAY)
    return save_diagram(image, "02-personal-flow.png")


def diagram_leader_simple():
    image, draw = new_canvas("领导发起转交员工流程", "领导交办可跳过普通审批，但不能跳过签收、反馈、验收评分和归档。", (1800, 900))
    draw.text((70, 185), "主线", font=font(28, True), fill=TEAL_DARK)
    draw.text((70, 560), "异常闭环", font=font(28, True), fill=RED)
    nodes = {}
    main = [
        ("领导发起\n督办", "a"),
        ("生成督办单\n或转交派发", "b"),
        ("派发\n责任人/时限", "c"),
        ("责任人\n签收", "d"),
        ("执行反馈", "e"),
        ("领导/督办员\n验收", "f"),
        ("评分归档", "g"),
    ]
    start_x, gap, w, h = 150, 42, 195, 102
    for idx, (text, key) in enumerate(main):
        nodes[key] = {
            "x": start_x + idx * (w + gap),
            "y": 250,
            "w": w,
            "h": h,
            "text": text,
            "font_size": 25,
            "fill": BLUE_LIGHT if key in {"d", "e", "f"} else TEAL_LIGHT,
            "outline": BLUE if key in {"d", "e", "f"} else TEAL,
        }
    abnormal = [
        ("未指定责任人\n转督办员/部门负责人", "r1", 240),
        ("责任人拒收\n退回派发并抄送领导", "r2", 610),
        ("超时未签收\n催签亮灯并升级", "r3", 980),
        ("验收退回\n整改后再验收", "r4", 1350),
    ]
    for text, key, x in abnormal:
        nodes[key] = {
            "x": x,
            "y": 620,
            "w": 275,
            "h": 98,
            "text": text,
            "font_size": 22,
            "fill": RED_LIGHT,
            "outline": RED,
        }
    for node in nodes.values():
        draw_node(draw, node)
    for left, right in zip(["a", "b", "c", "d", "e", "f"], ["b", "c", "d", "e", "f", "g"]):
        draw_edge(draw, nodes, left, right)
    draw_edge(draw, nodes, "b", "r1", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r1", "c", "top", "bottom", color=RED)
    draw_edge(draw, nodes, "d", "r2", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r2", "c", "top", "bottom", color=RED)
    draw_edge(draw, nodes, "d", "r3", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r3", "b", "top", "bottom", color=RED)
    draw_edge(draw, nodes, "f", "r4", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r4", "e", "top", "bottom", color=RED)
    draw.rounded_rectangle((112, 775, 1688, 835), radius=18, fill=GRAY_LIGHT, outline=(203, 213, 225), width=3)
    draw.text((145, 792), "控制点：领导交办要快速生成督办，但责任人拒收、超时未签收必须抄送和升级，避免指示落空。", font=font(24, True), fill=GRAY)
    return save_diagram(image, "03-leader-flow.png")


def diagram_department_simple():
    image, draw = new_canvas("部门发起督办流程", "部门督办重点处理本部门派发和跨部门承接确认；跨部门拒绝必须协调或批示。", (1800, 900))
    draw.text((70, 185), "主线", font=font(28, True), fill=TEAL_DARK)
    draw.text((70, 560), "异常闭环", font=font(28, True), fill=RED)
    nodes = {}
    main = [
        ("部门填写\n督办事项", "a"),
        ("部门负责人\n确认", "b"),
        ("生成部门\n督办单", "c"),
        ("本部门派发\n或跨部门确认", "d"),
        ("责任人\n签收", "e"),
        ("执行反馈", "f"),
        ("部门验收\n评分归档", "g"),
    ]
    start_x, gap, w, h = 150, 42, 195, 102
    for idx, (text, key) in enumerate(main):
        nodes[key] = {
            "x": start_x + idx * (w + gap),
            "y": 250,
            "w": w,
            "h": h,
            "text": text,
            "font_size": 24,
            "fill": BLUE_LIGHT if key in {"e", "f", "g"} else TEAL_LIGHT,
            "outline": BLUE if key in {"e", "f", "g"} else TEAL,
        }
    abnormal = [
        ("确认退回\n补充事项信息", "r1", 245),
        ("跨部门拒绝承接\n协调或领导批示", "r2", 625),
        ("拒收/签收超时\n催签、重派、升级", "r3", 1005),
        ("验收退回\n整改后再提交", "r4", 1385),
    ]
    for text, key, x in abnormal:
        nodes[key] = {
            "x": x,
            "y": 620,
            "w": 275,
            "h": 98,
            "text": text,
            "font_size": 22,
            "fill": RED_LIGHT,
            "outline": RED,
        }
    for node in nodes.values():
        draw_node(draw, node)
    for left, right in zip(["a", "b", "c", "d", "e", "f"], ["b", "c", "d", "e", "f", "g"]):
        draw_edge(draw, nodes, left, right)
    draw_edge(draw, nodes, "b", "r1", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r1", "a", "top", "bottom", color=RED)
    draw_edge(draw, nodes, "d", "r2", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r2", "d", "top", "bottom", color=RED)
    draw_edge(draw, nodes, "e", "r3", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r3", "d", "top", "bottom", color=RED)
    draw_edge(draw, nodes, "g", "r4", "bottom", "top", color=RED)
    draw_edge(draw, nodes, "r4", "f", "top", "bottom", color=RED)
    draw.rounded_rectangle((110, 775, 1690, 835), radius=18, fill=GRAY_LIGHT, outline=(203, 213, 225), width=3)
    draw.text((145, 792), "控制点：部门督办要先确认责任范围；跨部门拒绝不能直接关闭，必须进入协调、批示或重新确认。", font=font(24, True), fill=GRAY)
    return save_diagram(image, "04-department-flow.png")


def diagram_personal():
    image, draw = new_canvas("个人发起督办流程", "个人发起必须先审批，审批成立后才派发责任人。", (1800, 1040))
    nodes = {
        "a": {"x": 60, "y": 200, "w": 185, "h": 92, "text": "个人填写\n督办申请", "font_size": 25},
        "b": {"x": 315, "y": 180, "w": 170, "h": 132, "text": "信息\n完整？", "kind": "decision", "fill": AMBER_LIGHT, "outline": AMBER, "font_size": 24},
        "c": {"x": 560, "y": 200, "w": 170, "h": 92, "text": "提交审批", "font_size": 26},
        "d": {"x": 795, "y": 180, "w": 175, "h": 132, "text": "督办\n成立？", "kind": "decision", "fill": AMBER_LIGHT, "outline": AMBER, "font_size": 24},
        "e": {"x": 1035, "y": 200, "w": 160, "h": 92, "text": "待派发", "font_size": 26},
        "f": {"x": 1260, "y": 200, "w": 215, "h": 92, "text": "派发责任人\n协办人/时限", "font_size": 24},
        "g": {"x": 1535, "y": 200, "w": 200, "h": 92, "text": "责任人\n待签收", "font_size": 25},
        "h": {"x": 1535, "y": 475, "w": 200, "h": 92, "text": "执行反馈", "font_size": 26, "fill": BLUE_LIGHT, "outline": BLUE},
        "i": {"x": 1260, "y": 475, "w": 215, "h": 92, "text": "提交完成", "font_size": 26, "fill": BLUE_LIGHT, "outline": BLUE},
        "j": {"x": 1035, "y": 475, "w": 160, "h": 92, "text": "验收", "font_size": 26, "fill": BLUE_LIGHT, "outline": BLUE},
        "k": {"x": 795, "y": 475, "w": 175, "h": 92, "text": "评分归档", "font_size": 26, "fill": TEAL_LIGHT, "outline": TEAL},
        "x": {"x": 560, "y": 475, "w": 170, "h": 92, "text": "关闭并\n记录原因", "font_size": 24, "fill": RED_LIGHT, "outline": RED},
        "r": {"x": 1260, "y": 740, "w": 215, "h": 92, "text": "拒收/超时\n回到派发", "font_size": 24, "fill": RED_LIGHT, "outline": RED},
    }
    for node in nodes.values():
        draw_node(draw, node)
    draw_edge(draw, nodes, "a", "b")
    draw_edge(draw, nodes, "b", "c", text="完整")
    draw_edge(draw, nodes, "b", "a", "bottom", "bottom", text="补充", via=[(400, 380), (150, 380)])
    draw_edge(draw, nodes, "c", "d")
    draw_edge(draw, nodes, "d", "e", text="成立")
    draw_edge(draw, nodes, "d", "a", "top", "top", text="退回补充", via=[(882, 135), (150, 135)])
    draw_edge(draw, nodes, "d", "x", "bottom", "top", text="不成立", color=RED)
    draw_edge(draw, nodes, "e", "f")
    draw_edge(draw, nodes, "f", "g")
    draw_edge(draw, nodes, "g", "h", "bottom", "top", text="签收")
    draw_edge(draw, nodes, "h", "i", "left", "right")
    draw_edge(draw, nodes, "i", "j", "left", "right")
    draw_edge(draw, nodes, "j", "k", "left", "right", text="通过")
    draw_edge(draw, nodes, "j", "h", "top", "left", text="退回整改", via=[(1115, 410), (1500, 410)])
    draw_edge(draw, nodes, "g", "r", "bottom", "top", text="拒收/超时", color=RED)
    draw_edge(draw, nodes, "r", "f", "top", "bottom", color=RED)
    return save_diagram(image, "02-personal-flow.png")


def diagram_leader():
    image, draw = new_canvas("领导发起转交员工流程", "领导交办可跳过普通审批，但拒收和超时必须升级可见。", (1800, 1040))
    nodes = {
        "a": {"x": 70, "y": 210, "w": 180, "h": 90, "text": "领导发起\n督办", "font_size": 25},
        "b": {"x": 315, "y": 188, "w": 175, "h": 132, "text": "是否直接\n指定责任人", "kind": "decision", "fill": AMBER_LIGHT, "outline": AMBER, "font_size": 23},
        "c": {"x": 565, "y": 120, "w": 190, "h": 90, "text": "系统生成\n督办单", "font_size": 25},
        "d": {"x": 565, "y": 340, "w": 205, "h": 90, "text": "转秘书/督办员\n部门负责人", "font_size": 23},
        "e": {"x": 830, "y": 340, "w": 205, "h": 90, "text": "派发责任人\n协办人/时限", "font_size": 23},
        "f": {"x": 1100, "y": 210, "w": 185, "h": 90, "text": "责任人\n待签收", "font_size": 25},
        "g": {"x": 1355, "y": 188, "w": 170, "h": 132, "text": "签收\n结果", "kind": "decision", "fill": AMBER_LIGHT, "outline": AMBER, "font_size": 25},
        "h": {"x": 1565, "y": 210, "w": 170, "h": 90, "text": "执行中", "font_size": 26, "fill": BLUE_LIGHT, "outline": BLUE},
        "i": {"x": 1565, "y": 480, "w": 170, "h": 90, "text": "阶段/完成\n反馈", "font_size": 24, "fill": BLUE_LIGHT, "outline": BLUE},
        "j": {"x": 1355, "y": 480, "w": 170, "h": 90, "text": "领导/督办员\n验收", "font_size": 23, "fill": BLUE_LIGHT, "outline": BLUE},
        "k": {"x": 1100, "y": 480, "w": 185, "h": 90, "text": "评分归档", "font_size": 26, "fill": TEAL_LIGHT, "outline": TEAL},
        "r": {"x": 1110, "y": 725, "w": 220, "h": 90, "text": "拒收：退回派发\n并抄送领导", "font_size": 22, "fill": RED_LIGHT, "outline": RED},
        "t": {"x": 1370, "y": 725, "w": 220, "h": 90, "text": "超时：催签亮灯\n升级负责人", "font_size": 22, "fill": RED_LIGHT, "outline": RED},
    }
    for node in nodes.values():
        draw_node(draw, node)
    draw_edge(draw, nodes, "a", "b")
    draw_edge(draw, nodes, "b", "c", "right", "left", "是")
    draw_edge(draw, nodes, "b", "d", "bottom", "left", "否", via=[(402, 385)])
    draw_edge(draw, nodes, "d", "e")
    draw_edge(draw, nodes, "c", "f", "right", "left", via=[(900, 165), (900, 255)])
    draw_edge(draw, nodes, "e", "f", "right", "left", via=[(1060, 385), (1060, 255)])
    draw_edge(draw, nodes, "f", "g")
    draw_edge(draw, nodes, "g", "h", text="签收")
    draw_edge(draw, nodes, "h", "i", "bottom", "top")
    draw_edge(draw, nodes, "i", "j", "left", "right")
    draw_edge(draw, nodes, "j", "k", "left", "right", text="通过")
    draw_edge(draw, nodes, "j", "h", "top", "bottom", text="退回整改", via=[(1440, 430), (1650, 430)])
    draw_edge(draw, nodes, "g", "r", "bottom", "top", text="拒收", color=RED)
    draw_edge(draw, nodes, "g", "t", "bottom", "top", text="超时", color=RED)
    draw_edge(draw, nodes, "r", "e", "left", "bottom", color=RED, via=[(900, 770), (900, 450)])
    draw_edge(draw, nodes, "t", "f", "left", "bottom", color=RED, via=[(1190, 770), (1190, 320)])
    return save_diagram(image, "03-leader-flow.png")


def diagram_department():
    image, draw = new_canvas("部门发起督办流程", "部门督办重点处理本部门派发和跨部门承接确认。", (1800, 1080))
    nodes = {
        "a": {"x": 70, "y": 210, "w": 180, "h": 90, "text": "部门填写\n督办事项", "font_size": 25},
        "b": {"x": 315, "y": 210, "w": 190, "h": 90, "text": "部门负责人\n确认", "font_size": 24},
        "c": {"x": 570, "y": 190, "w": 170, "h": 132, "text": "是否\n通过", "kind": "decision", "fill": AMBER_LIGHT, "outline": AMBER, "font_size": 25},
        "d": {"x": 805, "y": 210, "w": 185, "h": 90, "text": "生成部门\n督办单", "font_size": 24},
        "e": {"x": 1055, "y": 190, "w": 175, "h": 132, "text": "责任\n范围", "kind": "decision", "fill": AMBER_LIGHT, "outline": AMBER, "font_size": 25},
        "f": {"x": 1295, "y": 120, "w": 190, "h": 90, "text": "本部门\n派发", "font_size": 25},
        "g": {"x": 1295, "y": 360, "w": 190, "h": 90, "text": "跨部门\n接收确认", "font_size": 24},
        "h": {"x": 1535, "y": 342, "w": 175, "h": 132, "text": "是否\n承接", "kind": "decision", "fill": AMBER_LIGHT, "outline": AMBER, "font_size": 25},
        "i": {"x": 1535, "y": 120, "w": 190, "h": 90, "text": "指定责任人", "font_size": 25},
        "j": {"x": 1295, "y": 610, "w": 190, "h": 90, "text": "协调/领导\n批示", "font_size": 24, "fill": RED_LIGHT, "outline": RED},
        "k": {"x": 1055, "y": 610, "w": 175, "h": 90, "text": "责任人\n签收", "font_size": 25, "fill": BLUE_LIGHT, "outline": BLUE},
        "l": {"x": 805, "y": 610, "w": 185, "h": 90, "text": "执行反馈", "font_size": 26, "fill": BLUE_LIGHT, "outline": BLUE},
        "m": {"x": 570, "y": 610, "w": 170, "h": 90, "text": "部门验收", "font_size": 26, "fill": BLUE_LIGHT, "outline": BLUE},
        "n": {"x": 315, "y": 610, "w": 190, "h": 90, "text": "评分归档", "font_size": 26, "fill": TEAL_LIGHT, "outline": TEAL},
    }
    for node in nodes.values():
        draw_node(draw, node)
    draw_edge(draw, nodes, "a", "b")
    draw_edge(draw, nodes, "b", "c")
    draw_edge(draw, nodes, "c", "d", text="通过")
    draw_edge(draw, nodes, "c", "a", "top", "top", text="退回", via=[(655, 145), (160, 145)])
    draw_edge(draw, nodes, "d", "e")
    draw_edge(draw, nodes, "e", "f", "right", "left", "本部门", via=[(1240, 255), (1240, 165)])
    draw_edge(draw, nodes, "e", "g", "bottom", "left", "跨部门", via=[(1142, 405)])
    draw_edge(draw, nodes, "g", "h")
    draw_edge(draw, nodes, "h", "i", "top", "bottom", text="同意")
    draw_edge(draw, nodes, "h", "j", "bottom", "top", text="拒绝", color=RED)
    draw_edge(draw, nodes, "j", "g", "top", "bottom", text="再协调", color=RED)
    draw_edge(draw, nodes, "f", "k", "bottom", "top", via=[(1390, 560), (1142, 560)])
    draw_edge(draw, nodes, "i", "k", "bottom", "top", via=[(1630, 560), (1142, 560)])
    draw_edge(draw, nodes, "k", "l", "left", "right")
    draw_edge(draw, nodes, "l", "m", "left", "right")
    draw_edge(draw, nodes, "m", "n", "left", "right", text="通过")
    draw_edge(draw, nodes, "m", "l", "top", "top", text="退回整改", via=[(655, 565), (900, 565)])
    return save_diagram(image, "04-department-flow.png")


def diagram_exception():
    image, draw = new_canvas("异常动作闭环流程", "延期、变更、转办、拒收、催办、批示都不能悬空，必须回到主线状态。", (1800, 1040))
    nodes = {
        "main": {"x": 760, "y": 445, "w": 280, "h": 110, "text": "督办主线\n待签收 / 执行中", "font_size": 28, "fill": TEAL_LIGHT, "outline": TEAL},
        "delay": {"x": 130, "y": 210, "w": 265, "h": 105, "text": "延期申请\n原因/新时限/影响", "font_size": 23, "fill": AMBER_LIGHT, "outline": AMBER},
        "change": {"x": 500, "y": 210, "w": 265, "h": 105, "text": "变更申请\n内容/原因/版本", "font_size": 23, "fill": AMBER_LIGHT, "outline": AMBER},
        "transfer": {"x": 1035, "y": 210, "w": 265, "h": 105, "text": "转办申请\n新责任人重新签收", "font_size": 23, "fill": AMBER_LIGHT, "outline": AMBER},
        "reject": {"x": 1405, "y": 210, "w": 265, "h": 105, "text": "拒收处理\n原因/建议责任人", "font_size": 23, "fill": RED_LIGHT, "outline": RED},
        "urge": {"x": 315, "y": 735, "w": 265, "h": 105, "text": "催办\n提醒并留痕统计", "font_size": 23, "fill": BLUE_LIGHT, "outline": BLUE},
        "instruct": {"x": 1220, "y": 735, "w": 265, "h": 105, "text": "领导批示\n形成新待办/状态变化", "font_size": 22, "fill": BLUE_LIGHT, "outline": BLUE},
    }
    for node in nodes.values():
        draw_node(draw, node)
    pairs = [
        ("delay", "审批同意/驳回"),
        ("change", "生成变更版本"),
        ("transfer", "同意后回待签收"),
        ("reject", "重派或驳回拒收"),
        ("urge", "不改变主线状态"),
        ("instruct", "推动状态变化"),
    ]
    for key, text in pairs:
        draw_edge(draw, nodes, "main", key, "top" if nodes[key]["y"] < 445 else "bottom", "bottom" if nodes[key]["y"] < 445 else "top", color=GRAY)
        draw_edge(draw, nodes, key, "main", "bottom" if nodes[key]["y"] < 445 else "top", "top" if nodes[key]["y"] < 445 else "bottom", text=text, color=TEAL)
    draw.rounded_rectangle((585, 890, 1215, 955), radius=18, fill=GRAY_LIGHT, outline=(203, 213, 225), width=3)
    draw.text((620, 908), "控制原则：每个异常动作必须有申请、审批/确认、处理结果、回到主线、过程留痕。", font=font(24, True), fill=GRAY)
    return save_diagram(image, "05-exception-loop.png")


def diagram_warning():
    image, draw = new_canvas("亮灯预警与升级流程", "灯色取最严重状态，领导看板优先展示红灯和黑灯事项。", (1800, 850))
    nodes = {
        "green": {"x": 100, "y": 250, "w": 240, "h": 110, "text": "绿灯\n正常反馈", "font_size": 27, "fill": (236, 253, 245), "outline": (22, 163, 74)},
        "yellow": {"x": 450, "y": 250, "w": 240, "h": 110, "text": "黄灯\n临期/即将超时", "font_size": 25, "fill": (254, 249, 195), "outline": (202, 138, 4)},
        "red": {"x": 800, "y": 250, "w": 240, "h": 110, "text": "红灯\n已超期/异常未处理", "font_size": 24, "fill": RED_LIGHT, "outline": RED},
        "black": {"x": 1150, "y": 250, "w": 240, "h": 110, "text": "黑灯\n严重超期/无人处理", "font_size": 24, "fill": (241, 245, 249), "outline": (15, 23, 42)},
        "leader": {"x": 1455, "y": 250, "w": 245, "h": 110, "text": "领导批示\n专项协调", "font_size": 25, "fill": BLUE_LIGHT, "outline": BLUE},
    }
    for node in nodes.values():
        draw_node(draw, node)
    draw_edge(draw, nodes, "green", "yellow", text="剩余时限 <20%")
    draw_edge(draw, nodes, "yellow", "red", text="超期或拒收")
    draw_edge(draw, nodes, "red", "black", text="严重超期")
    draw_edge(draw, nodes, "black", "leader", text="升级")
    actions = [
        (150, 515, "正常展示"),
        (500, 515, "提醒责任人"),
        (850, 515, "抄送上级/督办人"),
        (1210, 515, "进入黑灯清单"),
        (1510, 515, "批示后生成动作"),
    ]
    for x, y, text in actions:
        draw.rounded_rectangle((x, y, x + 210, y + 70), radius=16, fill=GRAY_LIGHT, outline=(203, 213, 225), width=3)
        draw_text_center(draw, (x, y, 210, 70), text, font(22, True), GRAY)
    return save_diagram(image, "06-warning-flow.png")


def build_charts():
    return {
        "overall": diagram_overall(),
        "personal": diagram_personal_simple(),
        "leader": diagram_leader_simple(),
        "department": diagram_department_simple(),
        "exception": diagram_exception(),
        "warning": diagram_warning(),
    }


def add_figure(doc, title, image_path, width=6.65):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_font(r, size=10.5, bold=True, color=(15, 83, 76))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))


def main():
    charts = build_charts()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[name].font.name = "Microsoft YaHei"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    cover(doc)

    heading(doc, "一、方案摘要", 1)
    callout(
        doc,
        "建设目标",
        "建立一套覆盖发起、审批/确认、派发、签收、执行反馈、异常处理、验收、评分、归档的督办闭环机制。重点解决事项派发后无人签收、责任人拒收、延期无审批、转办无闭环、催办无留痕、领导批示无法落地等管理痛点。",
    )
    body(
        doc,
        "本方案将督办事项作为独立的管理对象，不依赖某一个具体 OA 审批流程。无论事项来自个人、领导还是部门，只要进入督办体系，就必须具备责任人、完成时限、过程反馈、异常处理、亮灯预警、验收评分和归档记录。",
    )
    add_figure(doc, "图 1：督办总体流转主线", charts["overall"], width=6.85)

    heading(doc, "二、三类督办流程概览", 1)
    table(
        doc,
        ["流程类型", "适用场景", "关键控制点", "闭环要求"],
        [
            ["个人发起督办", "员工发现跨部门事项、长期未推动事项、本人事项需正式跟踪", "发起后必须审批；审批通过后才允许派发", "验收通过并评分后办结"],
            ["领导发起督办", "领导指示、会议纪要、重点专项、临时交办事项", "可直接生成督办；责任人必须签收；拒收需抄送领导", "领导或督办员验收评分后归档"],
            ["部门发起督办", "部门间协作、制度执行、项目推进、专项整改", "部门负责人确认；跨部门需接收确认", "跨部门拒绝进入协调或领导批示"],
        ],
        widths=[1.2, 2.05, 2.15, 1.75],
    )

    heading(doc, "三、个人发起督办流程", 1)
    body(doc, "个人发起必须先确认督办是否成立，避免普通协作事项被随意升级为正式督办。审批退回、不成立、拒收、超时未签收都必须有明确处理路径。")
    add_figure(doc, "图 2：个人发起督办流程图", charts["personal"], width=6.85)

    heading(doc, "四、领导发起转交员工流程", 1)
    body(doc, "领导交办可以跳过普通审批，但不能跳过责任人签收、执行反馈、验收评分和归档。拒收或超时未签收时，系统必须抄送领导或部门负责人。")
    add_figure(doc, "图 3：领导发起转交员工流程图", charts["leader"], width=6.85)

    heading(doc, "五、部门发起督办流程", 1)
    body(doc, "部门发起需要先由部门负责人确认。跨部门督办必须增加接收部门承接确认，拒绝承接时进入协调或领导批示，不能直接关闭。")
    add_figure(doc, "图 4：部门发起督办流程图", charts["department"], width=6.85)

    heading(doc, "六、统一主线节点规则", 1)
    table(
        doc,
        ["阶段", "核心动作", "责任角色", "关键规则"],
        [
            ["1. 发起", "填写督办事项、目标、时限、附件", "发起人 / 领导 / 部门", "信息不完整不能提交"],
            ["2. 审批或确认", "判断督办是否成立", "审批人 / 部门负责人", "不成立必须说明原因"],
            ["3. 派发", "指定责任人、协办人、完成时限", "派发人 / 督办员", "没有责任人和时限不能派发"],
            ["4. 签收", "责任人确认接收任务", "责任人", "不签收必须催签并升级"],
            ["5. 执行反馈", "提交进度、卡点、附件、完成结果", "责任人 / 协办人", "长期无反馈自动催办"],
            ["6. 异常处理", "延期、变更、转办、拒收、批示", "相关责任角色", "异常不能悬空，必须回到主线"],
            ["7. 验收", "确认完成质量和交付物", "发起人 / 督办人 / 领导", "退回必须写整改要求"],
            ["8. 评分", "按时效、质量、协同、材料完整性评分", "验收人 / 领导", "评分后才允许办结"],
            ["9. 归档", "形成结论和过程留痕", "督办人 / 系统", "归档后进入统计分析"],
        ],
        widths=[1.05, 2.05, 1.55, 2.5],
    )

    heading(doc, "七、关键闭环机制", 1)
    heading(doc, "1. 签收闭环", 2)
    body(doc, "派发后必须进入签收环节。责任人签收时需要确认任务范围、完成标准、完成时限和交付物。未签收、延期签收、拒收都必须进入明确处理路径。")
    table(
        doc,
        ["情形", "系统动作", "处理结果"],
        [
            ["正常签收", "记录签收时间和签收人", "进入执行中"],
            ["4 小时未签收", "提醒责任人", "保持待签收"],
            ["8 小时未签收", "抄送派发人并亮黄灯", "进入签收预警"],
            ["1 个工作日未签收", "抄送责任人上级并亮红灯", "进入签收超期"],
            ["仍未签收", "要求派发人改派、强制签收或升级批示", "进入升级处理中"],
            ["责任人拒收", "填写拒收原因和建议责任人", "派发人同意后重派，驳回后继续签收"],
        ],
        widths=[1.35, 3.05, 2.65],
    )

    heading(doc, "2. 异常闭环", 2)
    add_figure(doc, "图 5：异常动作闭环流程图", charts["exception"], width=6.85)
    table(
        doc,
        ["异常动作", "必须填写内容", "闭环规则"],
        [
            ["延期", "延期原因、新完成时间、影响范围", "同意后更新时限并记录延期次数；驳回则维持原时限"],
            ["变更", "变更内容、变更原因、影响说明", "生成变更版本，不能覆盖原督办要求"],
            ["转办", "转办原因、新责任人、责任说明", "新责任人必须重新签收"],
            ["拒收", "拒收原因、建议责任人、是否需协调", "不能直接关闭，必须回到派发或领导批示"],
            ["批示", "批示意见、承接动作、责任人", "批示必须形成新待办或推动状态变化"],
            ["催办", "催办对象、催办内容、催办时间", "必须留痕并进入统计"],
        ],
        widths=[1.1, 2.75, 3.2],
    )

    heading(doc, "八、亮灯预警规则", 1)
    body(doc, "亮灯规则用于让领导和督办人员快速识别风险。建议将签收、进度、到期、验收分别亮灯，总体灯色取最严重状态。")
    add_figure(doc, "图 6：亮灯预警与升级流程图", charts["warning"], width=6.85)
    table(
        doc,
        ["灯色", "触发条件", "系统动作"],
        [
            ["绿灯", "未临期，正常反馈", "正常展示"],
            ["黄灯", "距截止时间小于 20%，或签收/反馈即将超时", "提醒责任人"],
            ["红灯", "已超期，或拒收/延期/转办未处理", "抄送上级和督办人"],
            ["黑灯", "严重超期，或重大事项无人处理", "升级领导批示"],
        ],
        widths=[1.0, 3.15, 2.9],
    )

    heading(doc, "九、验收评分机制", 1)
    body(doc, "督办事项完成后不应直接关闭，必须先验收，再评分，最后归档。评分结果可用于部门执行力分析和责任人绩效参考。")
    table(
        doc,
        ["评分项", "分值", "说明"],
        [
            ["按时完成", "30", "是否按计划完成，是否发生超期"],
            ["完成质量", "30", "是否达到督办目标和交付标准"],
            ["反馈及时性", "20", "是否按要求反馈进展和卡点"],
            ["协同配合", "10", "是否主动协同相关部门"],
            ["材料完整性", "10", "附件、说明、过程记录是否完整"],
        ],
        widths=[1.4, 0.75, 4.9],
    )
    body(doc, "评分建议：90-100 为优秀，80-89 为良好，60-79 为合格，60 分以下为不合格并进入复盘。")

    heading(doc, "十、领导关注看板", 1)
    body(doc, "领导端不需要处理所有细节，应聚焦重大事项、超期事项、拒收事项和跨部门争议。建议建设以下看板指标。")
    bullets(
        doc,
        [
            "督办总量、待签收数量、执行中数量、已超期数量。",
            "按部门、责任人、灯色统计的风险分布。",
            "重大督办事项清单和黑灯事项清单。",
            "拒收、延期、转办、变更次数排行。",
            "平均办结周期、按时完成率、验收通过率、评分分布。",
        ],
    )

    heading(doc, "十一、落地建议", 1)
    table(
        doc,
        ["阶段", "建设重点", "交付结果"],
        [
            ["第一阶段", "统一督办单、状态、签收、催办、亮灯规则", "形成可运行的基础督办闭环"],
            ["第二阶段", "补齐延期、变更、转办、拒收、批示机制", "解决异常事项悬空问题"],
            ["第三阶段", "上线验收评分和领导看板", "形成执行力分析和管理抓手"],
            ["第四阶段", "接入会议纪要、项目计划、重点工作来源", "扩展为统一事项督办平台"],
        ],
        widths=[1.1, 3.3, 2.65],
    )

    callout(
        doc,
        "结论",
        "督办系统的关键不是“能发起”，而是“必须闭环”。建议将签收、拒收、延期、变更、转办、催办、批示、验收、评分全部设计为可追踪、可升级、可统计的流程动作，确保事项不会停在某个无人处理的状态。",
    )

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("督办全流程设计方案｜领导汇报版｜含流程图")
        set_font(run, size=9, color=(120, 130, 140))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    shutil.copyfile(OUTPUT, VERIFY_COPY)
    print(VERIFY_COPY)


if __name__ == "__main__":
    main()
