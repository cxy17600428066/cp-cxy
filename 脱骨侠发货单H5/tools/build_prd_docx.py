from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "脱骨侠物流履约系统PRD_V1.0.md"
OUTPUT = ROOT / "脱骨侠物流履约系统PRD_V1.0.docx"
SKILL_ROOT = Path(r"C:\Users\admin\.codex\plugins\cache\openai-primary-runtime\documents\26.805.11740\skills\documents")
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(36, 29, 27)
MUTED = RGBColor(104, 95, 91)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
PALE_RED = "FFF4F1"
RED = RGBColor(185, 31, 23)
WHITE = RGBColor(255, 255, 255)
CONTENT_WIDTH_DXA = 9360


def set_run_font(run, size=None, color=None, bold=None, italic=None, east_asia="Microsoft YaHei"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, *, bold=False, color=INK, size=9.2, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def paragraph_bottom_border(paragraph, color="B91F17", size="12", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    field_run._r.append(fld_char1)
    field_run._r.append(instr_text)
    field_run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=MUTED)


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.clear()
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    left = hp.add_run("脱骨侠物流履约系统 PRD")
    set_run_font(left, size=9, color=MUTED, bold=True)
    right = hp.add_run("\tV1.0 研发评审稿")
    set_run_font(right, size=9, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    add_page_field(fp)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def create_numbering(doc, num_fmt, lvl_text, *, bullet_font=None):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_fmt)
    lvl.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), lvl_text)
    lvl.append(text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    if bullet_font:
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), bullet_font)
        r_fonts.set(qn("w:hAnsi"), bullet_font)
        r_pr.append(r_fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_paragraph(doc, text, num_id):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    set_run_font(r, size=10, color=RED, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("脱骨侠物流履约及发货单二维码系统")
    set_run_font(r, size=25, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("业务流程、移动端 H5、PC 管理后台及全链路数据规则")
    set_run_font(r, size=13.5, color=MUTED)

    meta = [
        ("版本", "V1.0（研发评审稿）"),
        ("日期", "2026-08-10"),
        ("适用端", "移动端 H5 / PC 管理后台"),
        ("评审角色", "产品、UI、前端、后端、测试、实施、业务运营"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(f"{label}：")
        set_run_font(r1, size=10.5, color=INK, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, color=INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(18)
    paragraph_bottom_border(rule)

    callout = doc.add_paragraph()
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    callout.paragraph_format.space_before = Pt(6)
    callout.paragraph_format.space_after = Pt(18)
    callout.paragraph_format.line_spacing = 1.18
    paragraph_shading(callout, PALE_RED)
    r1 = callout.add_run("文档定位\n")
    set_run_font(r1, size=11, color=RED, bold=True)
    r2 = callout.add_run("以当前交互原型为基线，明确正式系统的业务对象、角色权限、状态流转、节点字段、后台能力、异常更正和验收标准，可直接用于研发方案评审与测试用例设计。")
    set_run_font(r2, size=10.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("核心业务闭环")
    set_run_font(r, size=11, color=DARK_BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run("平台订单  →  创建发货单  →  工厂发货  →  司机揽收  →  司机到货  →  甲方确认  →  履约完成")
    set_run_font(r, size=11.5, color=INK, bold=True)
    doc.add_page_break()


def clean_inline(text):
    return text.replace("`", "").replace("**", "")


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [clean_inline(x.strip()) for x in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", x) for x in cells):
            rows.append(cells)
        i += 1
    return rows, i


def table_widths(rows):
    cols = len(rows[0])
    max_lens = []
    for c in range(cols):
        length = max(len(row[c]) if c < len(row) else 0 for row in rows)
        max_lens.append(min(max(length, 7), 38))
    if cols >= 3:
        max_lens[0] = min(max_lens[0], 15)
    return column_widths_from_weights(max_lens, CONTENT_WIDTH_DXA)


def add_table(doc, rows):
    if not rows:
        return
    cols = len(rows[0])
    normalized = [row + [""] * (cols - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(normalized):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
                set_cell_text(cell, text, bold=True, color=DARK_BLUE, size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                is_short = len(text) <= 12 and cols > 2
                set_cell_text(cell, text, size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER if is_short else WD_ALIGN_PARAGRAPH.LEFT)
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(
        table,
        table_widths(normalized),
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=120,
        cell_margins_dxa={"top": 100, "bottom": 100, "start": 120, "end": 120},
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(clean_inline(text))
    set_run_font(run, size=11, color=INK)
    return p


def add_callout(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.14)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    paragraph_shading(p, LIGHT_BLUE)
    r = p.add_run(clean_inline(text))
    set_run_font(r, size=10.3, color=DARK_BLUE, bold=True)


def build():
    doc = Document()
    setup_page(doc)
    setup_styles(doc)
    doc.core_properties.title = "脱骨侠物流履约及发货单二维码系统 PRD"
    doc.core_properties.subject = "业务流程、H5、PC后台、数据与验收要求"
    doc.core_properties.author = "脱骨侠产品团队"
    doc.core_properties.keywords = "脱骨侠, 物流履约, 发货单, 二维码, H5, PRD"
    add_cover(doc)

    bullet_num = create_numbering(doc, "bullet", "•", bullet_font="Symbol")
    decimal_num = None

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1."))
    i = start
    pending = []

    def flush_pending():
        nonlocal pending
        if pending:
            add_body_paragraph(doc, " ".join(x.strip() for x in pending))
            pending = []

    page_break_sections = {"9. PC 管理后台详细需求", "18. 评审结论记录"}
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            flush_pending()
            decimal_num = None
            i += 1
            continue
        if stripped.startswith("|"):
            flush_pending()
            decimal_num = None
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if stripped.startswith("#### "):
            flush_pending()
            decimal_num = None
            p = doc.add_paragraph(clean_inline(stripped[5:]), style="Heading 3")
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        if stripped.startswith("### "):
            flush_pending()
            decimal_num = None
            title = clean_inline(stripped[4:])
            p = doc.add_paragraph(title, style="Heading 2")
            if title == "6.2 PC 管理后台":
                p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        if stripped.startswith("## "):
            flush_pending()
            decimal_num = None
            title = clean_inline(stripped[3:])
            p = doc.add_paragraph(title, style="Heading 1")
            if any(title.startswith(x) for x in page_break_sections):
                p.paragraph_format.page_break_before = True
            i += 1
            continue
        if stripped.startswith("> "):
            flush_pending()
            decimal_num = None
            add_callout(doc, stripped[2:])
            i += 1
            continue
        if stripped.startswith("- "):
            flush_pending()
            decimal_num = None
            add_list_paragraph(doc, clean_inline(stripped[2:]), bullet_num)
            i += 1
            continue
        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            flush_pending()
            if decimal_num is None:
                decimal_num = create_numbering(doc, "decimal", "%1.")
            add_list_paragraph(doc, clean_inline(m.group(1)), decimal_num)
            i += 1
            continue
        decimal_num = None
        pending.append(stripped)
        i += 1
    flush_pending()

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
