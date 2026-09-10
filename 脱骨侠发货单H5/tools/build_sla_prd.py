from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'订单履约时效管理PRD_V1.0_20260910.md'
OUTPUT=ROOT/'订单履约时效管理PRD_V1.0_20260910.docx'
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.72);section.bottom_margin=Inches(.68)
section.left_margin=Inches(.85);section.right_margin=Inches(.85)
section.header_distance=Inches(.25);section.footer_distance=Inches(.3)

def font(style,size,bold=False,color='000000'):
    style.font.name='Microsoft YaHei'
    style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor.from_string(color)
    pr=style.element.get_or_add_rPr()
    rf=pr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');pr.append(rf)
    for key in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+key),'Microsoft YaHei')
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'zh-CN');lang.set(qn('w:eastAsia'),'zh-CN');pr.append(lang)
font(doc.styles['Normal'],11)
doc.styles['Normal'].paragraph_format.line_spacing=1.22
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
font(doc.styles['Title'],24,True)
doc.styles['Title'].paragraph_format.space_after=Pt(15)
for style,size in [('Heading 1',18),('Heading 2',13),('Heading 3',11.5)]:
    font(doc.styles[style],size,True)
    pf=doc.styles[style].paragraph_format
    pf.space_before=Pt(12);pf.space_after=Pt(7);pf.keep_with_next=True
font(doc.styles['List Bullet'],11)
doc.styles['List Bullet'].paragraph_format.space_after=Pt(5)
doc.styles['List Bullet'].paragraph_format.line_spacing=1.2
for name in ['Title','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name]
    for b in style.element.xpath('.//w:pBdr'):b.getparent().remove(b)
    for c in style.element.xpath('.//w:color'):
        for attr in ('themeColor','themeTint','themeShade'):c.attrib.pop(qn('w:'+attr),None)
        c.set(qn('w:val'),'000000')

foot=section.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for text in ['第 ',None,' 页']:
    run=foot.add_run(text or '')
    run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string('666666')
    if text is None:
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
        run._r.addnext(field)

def widths(headers):
    if headers==['节点','起算事件','完成事件','统计说明']:return [1.05,1.55,2.55,1.65]
    if headers[0]=='编号':return [.58,2.56,3.66]
    if len(headers)==4:return [1.15,1.4,1.2,3.05]
    if headers[0]=='事件或数据':return [2.2,2.0,2.6]
    if headers[0]=='菜单分组':return [1.05,1.35,4.4]
    if len(headers)==3:return [1.45,2.0,3.35]
    return [1.72,5.08]

def table(rows):
    table=doc.add_table(rows=1,cols=len(rows[0]))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    dimensions=widths(rows[0])
    for col,w in zip(table.columns,dimensions):col.width=Inches(w)
    tblPr=table._tbl.tblPr
    layout=tblPr.find(qn('w:tblLayout'))
    if layout is None:layout=OxmlElement('w:tblLayout');tblPr.append(layout)
    layout.set(qn('w:type'),'fixed')
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    tblPr.append(borders)
    for i,values in enumerate(rows):
        row=table.rows[0] if i==0 else table.add_row()
        props=row._tr.get_or_add_trPr()
        noSplit=OxmlElement('w:cantSplit');props.append(noSplit)
        if i==0:props.append(OxmlElement('w:tblHeader'))
        for j,(cell,value,w) in enumerate(zip(row.cells,values,dimensions)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'253F5D' if i==0 else ('F4F7FA' if i%2==0 else 'FFFFFF'));pr.append(shading)
            margins=OxmlElement('w:tcMar')
            for edge,val in [('top','80'),('bottom','80'),('left','105'),('right','105')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.15
            p.paragraph_format.keep_with_next=i==0
            if j==0 or (i==0):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            run=p.add_run(value.strip())
            run.font.size=Pt(10.5);run.font.bold=i==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.line_spacing=1
    p.add_run('').font.size=Pt(2)

lines=SOURCE.read_text(encoding='utf-8').splitlines()
idx=0
while idx<len(lines):
    line=lines[idx].strip()
    if not line:idx+=1;continue
    if line=='<!-- pagebreak -->':
        doc.add_page_break();idx+=1;continue
    if line.startswith('|'):
        data=[]
        while idx<len(lines) and lines[idx].strip().startswith('|'):
            row=[v.strip() for v in lines[idx].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',v) for v in row):data.append(row)
            idx+=1
        table(data);continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:],style='Heading 2')
    elif line.startswith('- '):
        doc.add_paragraph(line[2:],style='List Bullet')
    else:
        p=doc.add_paragraph(line)
        if line.startswith(('版本 V','评审对象','文档状态')):
            p.paragraph_format.space_after=Pt(4)
            for run in p.runs:run.font.size=Pt(10);run.font.color.rgb=RGBColor.from_string('555555')
    idx+=1
doc.core_properties.title='订单履约时效管理产品需求文档'
doc.core_properties.subject='订单履约时效管理业务需求与验收'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.version='1.0'
doc.core_properties.comments=''
doc.save(OUTPUT)
print(OUTPUT)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
